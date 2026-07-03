import json
import io
import importlib
import shutil
import tempfile
import uuid
from contextlib import nullcontext
from datetime import date, datetime, time, timedelta
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.core.signing import SignatureExpired
from django.contrib.auth.models import AnonymousUser, Permission, User
from django.contrib.sessions.middleware import SessionMiddleware
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from django.urls import resolve, reverse
from django.urls.exceptions import Resolver404
from openpyxl import load_workbook
import pandas as pd

from attendance.forms import AttendanceActivityExportForm, AttendanceExportForm
from attendance.export_jobs import create_export_job, file_path, update_job
from attendance.methods import utils as attendance_utils
from attendance.models import (
    Attendance,
    AttendanceActivity,
    AttendanceLateComeEarlyOut,
    AttendanceOverTime,
)
from attendance.views import clock_in_out as clock_in_out_views
from attendance.views.clock_in_out import clock_out_attendance_and_activity
from attendance.views.portal import (
    _activity_duration_seconds,
    _geofence_check,
    _get_portal_pin_reset_employee,
    _maybe_auto_checkout_employee,
    _make_portal_pin_reset_token,
    _portal_activity_total_seconds,
    _portal_break_policy,
    _portal_worked_duration_metadata,
    _send_portal_pin_reset_email,
    attendance_history,
    employee_lookup,
    forgot_pin,
    public_activity_transition,
    public_clock_in,
    public_clock_out,
    reset_pin,
    verify_pin,
)
from attendance.views import views as attendance_views
from attendance.views.views import (
    _attendance_activity_daily_export_value,
    _attendance_activity_export_columns,
    _attendance_activity_export_data,
    _attendance_activity_export_data_from_rows,
    _attendance_activity_payroll_group_daily_rows,
    _build_export_total_ranges,
    _delete_blocked_message,
    _insert_export_employee_separator_rows,
    _payroll_group_export_selected_columns,
    _write_export_row_formulas,
    _write_export_totals_sheet,
    build_attendance_tab_context,
    build_daily_activity_rows,
    build_my_attendance_activity_meta,
)
from base.models import EmployeeShift, EmployeeShiftDay, EmployeeShiftSchedule, WorkType
from employee.models import Employee
from horilla.horilla_middlewares import _thread_locals


def attach_session(request):
    middleware = SessionMiddleware(lambda req: HttpResponse())
    middleware.process_request(request)
    request.user = getattr(request, "user", AnonymousUser())
    return request


class ClockOutAttendanceAndActivityTests(SimpleTestCase):
    def _setup_clock_out_mocks(self):
        employee = MagicMock()
        today = date(2026, 4, 10)
        out_dt = datetime(2026, 4, 10, 17, 0, 0)

        attendance_activity = MagicMock()
        attendance_activity.attendance_date = today

        first_activity = MagicMock()
        second_activity = MagicMock()

        activities_qs = MagicMock()
        open_qs = MagicMock()
        open_qs.exists.return_value = True
        open_qs.last.return_value = attendance_activity

        activities_qs.filter.side_effect = lambda **kwargs: (
            open_qs if kwargs == {"clock_out__isnull": True} else [first_activity, second_activity]
        )

        attendance = MagicMock()
        attendance_qs = MagicMock()
        attendance_qs.order_by.return_value = [attendance]

        return {
            "employee": employee,
            "today": today,
            "out_dt": out_dt,
            "attendance": attendance,
            "activities_qs": activities_qs,
            "attendance_qs": attendance_qs,
            "first_activity": first_activity,
            "second_activity": second_activity,
        }

    @patch("attendance.views.clock_in_out.attendance_validate", return_value=True)
    @patch("attendance.views.clock_in_out.overtime_calculation", return_value="00:00")
    @patch("attendance.views.clock_in_out.calculate_worked_hours", return_value="08:00")
    @patch("attendance.views.clock_in_out.Attendance")
    @patch("attendance.views.clock_in_out.AttendanceActivity")
    def test_clock_out_auto_validate_true_keeps_existing_behavior(
        self,
        attendance_activity_model,
        attendance_model,
        _calculate_worked_hours_mock,
        _overtime_mock,
        attendance_validate_mock,
    ):
        ctx = self._setup_clock_out_mocks()
        attendance_activity_model.objects.filter.return_value.order_by.return_value = ctx["activities_qs"]
        attendance_model.objects.filter.return_value = ctx["attendance_qs"]

        result = clock_out_attendance_and_activity(
            employee=ctx["employee"],
            date_today=ctx["today"],
            now="17:00",
            out_datetime=ctx["out_dt"],
            auto_validate=True,
        )

        self.assertEqual(result, ctx["attendance"])
        attendance_validate_mock.assert_called_once_with(ctx["attendance"])
        self.assertTrue(ctx["attendance"].attendance_validated)
        self.assertFalse(ctx["attendance"]._skip_auto_approve_overtime)

    @patch("attendance.views.clock_in_out.attendance_validate")
    @patch("attendance.views.clock_in_out.overtime_calculation", return_value="00:00")
    @patch("attendance.views.clock_in_out.calculate_worked_hours", return_value="08:00")
    @patch("attendance.views.clock_in_out.Attendance")
    @patch("attendance.views.clock_in_out.AttendanceActivity")
    def test_clock_out_auto_validate_false_forces_not_validated(
        self,
        attendance_activity_model,
        attendance_model,
        _calculate_worked_hours_mock,
        _overtime_mock,
        attendance_validate_mock,
    ):
        ctx = self._setup_clock_out_mocks()
        attendance_activity_model.objects.filter.return_value.order_by.return_value = ctx["activities_qs"]
        attendance_model.objects.filter.return_value = ctx["attendance_qs"]

        result = clock_out_attendance_and_activity(
            employee=ctx["employee"],
            date_today=ctx["today"],
            now="17:00",
            out_datetime=ctx["out_dt"],
            auto_validate=False,
        )

        self.assertEqual(result, ctx["attendance"])
        attendance_validate_mock.assert_not_called()
        self.assertFalse(ctx["attendance"].attendance_validated)
        self.assertFalse(ctx["attendance"]._skip_auto_approve_overtime)

    @patch("attendance.views.clock_in_out.attendance_validate", return_value=True)
    @patch("attendance.views.clock_in_out.overtime_calculation", return_value="01:00")
    @patch("attendance.views.clock_in_out.calculate_worked_hours", return_value="09:00")
    @patch("attendance.views.clock_in_out.Attendance")
    @patch("attendance.views.clock_in_out.AttendanceActivity")
    def test_clock_out_can_skip_overtime_auto_approval(
        self,
        attendance_activity_model,
        attendance_model,
        _calculate_worked_hours_mock,
        _overtime_mock,
        _attendance_validate_mock,
    ):
        ctx = self._setup_clock_out_mocks()
        attendance_activity_model.objects.filter.return_value.order_by.return_value = ctx["activities_qs"]
        attendance_model.objects.filter.return_value = ctx["attendance_qs"]

        result = clock_out_attendance_and_activity(
            employee=ctx["employee"],
            date_today=ctx["today"],
            now="17:00",
            out_datetime=ctx["out_dt"],
            auto_validate=True,
            auto_approve_overtime=False,
        )

        self.assertEqual(result, ctx["attendance"])
        self.assertTrue(ctx["attendance"].attendance_validated)
        self.assertTrue(ctx["attendance"]._skip_auto_approve_overtime)


class AttendanceOvertimeAutoApprovalTests(SimpleTestCase):
    def _attendance_with_overtime(self):
        attendance = Attendance()
        attendance.is_validate_request = False
        attendance.overtime_second = attendance_utils.strtime_seconds("01:00")
        attendance.attendance_overtime_approve = False
        return attendance

    def _auto_approval_condition(self):
        return SimpleNamespace(
            auto_approve_ot=True,
            minimum_overtime_to_approve="00:30",
            overtime_cutoff=None,
        )

    @patch("attendance.models.AttendanceValidationCondition.objects.first")
    def test_overtime_auto_approval_runs_by_default(self, condition_mock):
        attendance = self._attendance_with_overtime()
        condition_mock.return_value = self._auto_approval_condition()

        attendance.handle_overtime_conditions()

        self.assertTrue(attendance.attendance_overtime_approve)

    @patch("attendance.models.AttendanceValidationCondition.objects.first")
    def test_overtime_auto_approval_can_be_skipped_per_save(self, condition_mock):
        attendance = self._attendance_with_overtime()
        attendance._skip_auto_approve_overtime = True
        condition_mock.return_value = self._auto_approval_condition()

        attendance.handle_overtime_conditions()

        self.assertFalse(attendance.attendance_overtime_approve)


class SharedAttendanceRecalculationTests(SimpleTestCase):
    class FakeAttendanceQuerySet(list):
        def select_related(self, *args):
            return self

        def iterator(self):
            return iter(self)

    class FakeActivityQuerySet(list):
        def order_by(self, *args):
            return self

        def first(self):
            return self[0] if self else None

        def last(self):
            return self[-1] if self else None

        def exists(self):
            return bool(self)

        def filter(self, **kwargs):
            if kwargs == {"clock_out__isnull": False}:
                return SharedAttendanceRecalculationTests.FakeActivityQuerySet(
                    [activity for activity in self if activity.clock_out is not None]
                )
            return self

    @patch("attendance.models.AttendanceActivity")
    def test_calculate_worked_hours_uses_closed_work_activities_only(
        self, activity_model
    ):
        employee = SimpleNamespace(id=1)
        attendance_date = date(2026, 4, 10)
        activity_model.objects.filter.return_value = [
            SimpleNamespace(
                clock_in_date=attendance_date,
                clock_in=time(8, 0),
                clock_out_date=attendance_date,
                clock_out=time(12, 0),
            ),
            SimpleNamespace(
                clock_in_date=attendance_date,
                clock_in=time(13, 0),
                clock_out_date=attendance_date,
                clock_out=time(17, 0),
            ),
        ]

        worked_hours = attendance_utils.calculate_worked_hours(
            employee, attendance_date
        )

        self.assertEqual(worked_hours, "08:00")
        activity_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date=attendance_date,
            activity_type="work",
            clock_out__isnull=False,
        )

    @patch("attendance.methods.utils.attendance_day_checking", return_value="08:00")
    @patch("attendance.methods.utils.shift_schedule_with_weekday_fallback")
    @patch("base.context_processors.enable_late_come_early_out_tracking")
    @patch("attendance.models.AttendanceLateComeEarlyOut")
    @patch("attendance.models.AttendanceActivity")
    @patch("attendance.models.Attendance")
    def test_recalculate_attendance_for_shift_refreshes_hours_and_late_early(
        self,
        attendance_model,
        activity_model,
        late_early_model,
        tracking_setting,
        schedule_fallback,
        _attendance_day_checking,
    ):
        shift = SimpleNamespace(id=1)
        employee = SimpleNamespace(id=10)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            employee_id=employee,
            attendance_date=attendance_date,
            attendance_day=SimpleNamespace(day="friday"),
            shift_id=shift,
            attendance_clock_in=None,
            attendance_clock_in_date=None,
            attendance_clock_out=None,
            attendance_clock_out_date=None,
            minimum_hour="00:00",
            attendance_worked_hour="00:00",
            attendance_validated=True,
            attendance_overtime_approve=False,
            attendance_overtime="00:00",
            at_work_second=None,
            overtime_second=None,
            approved_overtime_second=0,
            save=MagicMock(),
        )
        work_activities = self.FakeActivityQuerySet(
            [
                SimpleNamespace(
                    clock_in_date=attendance_date,
                    clock_in=time(9, 30),
                    clock_out_date=attendance_date,
                    clock_out=time(12, 0),
                ),
                SimpleNamespace(
                    clock_in_date=attendance_date,
                    clock_in=time(13, 0),
                    clock_out_date=attendance_date,
                    clock_out=time(16, 0),
                ),
            ]
        )
        duration_activities = [
            SimpleNamespace(
                clock_in_date=attendance_date,
                clock_in=time(9, 30),
                clock_out_date=attendance_date,
                clock_out=time(12, 0),
            ),
            SimpleNamespace(
                clock_in_date=attendance_date,
                clock_in=time(13, 0),
                clock_out_date=attendance_date,
                clock_out=time(16, 0),
            ),
        ]
        schedule = SimpleNamespace(
            minimum_working_hour="08:00",
            start_time=time(8, 0),
            end_time=time(17, 0),
        )
        attendance_model.objects.filter.return_value = self.FakeAttendanceQuerySet(
            [attendance]
        )
        schedule_fallback.return_value = schedule
        tracking_setting.return_value = {"tracking": True}
        delete_qs = MagicMock()
        late_early_model.objects.filter.return_value = delete_qs

        def activity_filter(**kwargs):
            if kwargs.get("clock_out__isnull") is False:
                return duration_activities
            return work_activities

        activity_model.objects.filter.side_effect = activity_filter

        count = attendance_utils.recalculate_attendance_for_shift(shift)

        self.assertEqual(count, 1)
        self.assertEqual(attendance.minimum_hour, "08:00")
        self.assertEqual(attendance.attendance_worked_hour, "05:30")
        self.assertEqual(attendance.attendance_clock_in, time(9, 30))
        self.assertEqual(attendance.attendance_clock_out, time(16, 0))
        self.assertTrue(attendance.attendance_validated)
        self.assertEqual(attendance.attendance_overtime, "00:00")
        self.assertEqual(attendance.at_work_second, 19800)
        self.assertEqual(attendance.overtime_second, 0)
        delete_qs.delete.assert_called_once()
        created_types = {
            call.kwargs["type"]
            for call in late_early_model.objects.get_or_create.call_args_list
        }
        self.assertEqual(created_types, {"late_come", "early_out"})
        attendance.save.assert_called_once()

    @patch("attendance.methods.utils.attendance_day_checking", return_value="06:00")
    @patch("attendance.methods.utils.shift_schedule_with_weekday_fallback")
    @patch("base.context_processors.enable_late_come_early_out_tracking")
    @patch("attendance.models.AttendanceLateComeEarlyOut")
    @patch("attendance.models.AttendanceActivity")
    @patch("attendance.models.Attendance")
    def test_recalculate_attendance_without_work_activities_preserves_worked_hours(
        self,
        attendance_model,
        activity_model,
        _late_early_model,
        tracking_setting,
        schedule_fallback,
        _attendance_day_checking,
    ):
        shift = SimpleNamespace(id=1)
        attendance = SimpleNamespace(
            employee_id=SimpleNamespace(id=10),
            attendance_date=date(2026, 4, 10),
            attendance_day=SimpleNamespace(day="friday"),
            shift_id=shift,
            attendance_clock_in=time(8, 0),
            attendance_clock_in_date=date(2026, 4, 10),
            attendance_clock_out=time(17, 0),
            attendance_clock_out_date=date(2026, 4, 10),
            minimum_hour="08:00",
            attendance_worked_hour="09:00",
            attendance_validated=True,
            attendance_overtime_approve=True,
            attendance_overtime="01:00",
            at_work_second=32400,
            overtime_second=3600,
            approved_overtime_second=3600,
            save=MagicMock(),
        )
        attendance_model.objects.filter.return_value = self.FakeAttendanceQuerySet(
            [attendance]
        )
        activity_model.objects.filter.return_value = self.FakeActivityQuerySet([])
        schedule_fallback.return_value = SimpleNamespace(
            minimum_working_hour="06:00",
            start_time=time(8, 0),
            end_time=time(14, 0),
        )
        tracking_setting.return_value = {"tracking": False}

        count = attendance_utils.recalculate_attendance_for_shift(shift)

        self.assertEqual(count, 1)
        self.assertEqual(attendance.attendance_worked_hour, "09:00")
        self.assertEqual(attendance.minimum_hour, "06:00")
        self.assertEqual(attendance.attendance_overtime, "00:00")
        self.assertTrue(attendance.attendance_validated)
        self.assertTrue(attendance.attendance_overtime_approve)
        self.assertEqual(attendance.approved_overtime_second, 0)
        attendance.save.assert_called_once()


class ShiftScheduleAttendanceRecalculationModelTests(TestCase):
    def setUp(self):
        self.shift_day, _created = EmployeeShiftDay.objects.get_or_create(day="monday")
        self.shift = EmployeeShift.objects.create(
            employee_shift="Recalc Regular",
            weekly_full_time="40:00",
            full_time="200:00",
        )
        self.schedule = EmployeeShiftSchedule.objects.create(
            day=self.shift_day,
            shift_id=self.shift,
            minimum_working_hour="06:00",
            start_time=time(8, 0),
            end_time=time(14, 0),
        )
        self.employee = Employee.objects.create(
            employee_first_name="Recalc",
            employee_last_name="User",
            email="recalc-user@example.com",
            phone="09170000001",
            gender="male",
            is_active=True,
        )

    @patch("base.context_processors.enable_late_come_early_out_tracking")
    def test_recalculate_rebuilds_approved_overtime_account_with_null_old_seconds(
        self, tracking_setting
    ):
        tracking_setting.return_value = {"tracking": False}
        attendance_date = date(2026, 6, 1)
        attendance = Attendance.objects.create(
            employee_id=self.employee,
            attendance_date=attendance_date,
            shift_id=self.shift,
            attendance_day=self.shift_day,
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            attendance_worked_hour="09:00",
            minimum_hour="08:00",
            attendance_validated=True,
            attendance_overtime_approve=True,
        )
        Attendance.objects.filter(pk=attendance.pk).update(
            at_work_second=None,
            overtime_second=None,
            approved_overtime_second=3600,
            attendance_overtime="01:00",
        )
        AttendanceActivity.objects.create(
            employee_id=self.employee,
            attendance_date=attendance_date,
            shift_day=self.shift_day,
            clock_in_date=attendance_date,
            clock_in=time(9, 0),
            clock_out_date=attendance_date,
            clock_out=time(18, 0),
            activity_type="work",
        )

        count = attendance_utils.recalculate_attendance_for_shift(self.shift)

        self.assertEqual(count, 1)
        attendance.refresh_from_db()
        self.assertTrue(attendance.attendance_validated)
        self.assertTrue(attendance.attendance_overtime_approve)
        self.assertEqual(attendance.minimum_hour, "06:00")
        self.assertEqual(attendance.attendance_worked_hour, "09:00")
        self.assertEqual(attendance.attendance_overtime, "04:00")
        self.assertEqual(attendance.at_work_second, 32400)
        self.assertEqual(attendance.overtime_second, 14400)
        self.assertEqual(attendance.approved_overtime_second, 14400)
        overtime_account = AttendanceOverTime.objects.get(
            employee_id=self.employee,
            month="june",
            year=2026,
        )
        self.assertEqual(overtime_account.worked_hours, "06:00")
        self.assertEqual(overtime_account.pending_hours, "00:00")
        self.assertEqual(overtime_account.overtime, "04:00")
        self.assertEqual(overtime_account.overtime_second, 14400)


class LateComeEarlyOutScheduleTests(SimpleTestCase):
    @patch("attendance.methods.utils.EmployeeShiftSchedule")
    def test_shift_schedule_today_uses_weekday_fallback(self, schedule_model):
        saturday = EmployeeShiftDay(id=6, day="saturday")
        shift = EmployeeShift(id=1, employee_shift="Morning")
        exact_qs = MagicMock()
        exact_qs.first.return_value = None
        fallback_qs = MagicMock()
        fallback_qs.first.return_value = SimpleNamespace(
            minimum_working_hour="08:00",
            start_time=time(8, 0),
            end_time=time(17, 0),
        )
        schedule_model.objects.filter.side_effect = [exact_qs, fallback_qs]

        self.assertEqual(
            attendance_utils.shift_schedule_today(saturday, shift),
            ("08:00", 28800, 61200),
        )

    @patch("attendance.views.clock_in_out.GraceTime")
    @patch("attendance.views.clock_in_out.enable_late_come_early_out_tracking")
    @patch("attendance.views.clock_in_out.late_come_create")
    def test_late_come_uses_attendance_clock_time(
        self, late_come_create, tracking, grace_time
    ):
        tracking.return_value = {"tracking": True}
        grace_time.objects.filter.return_value.exists.return_value = False
        attendance = SimpleNamespace(attendance_clock_in=time(10, 0))

        clock_in_out_views.late_come(
            attendance=attendance,
            start_time=attendance_utils.strtime_seconds("09:00"),
            end_time=attendance_utils.strtime_seconds("17:00"),
            shift=None,
        )

        late_come_create.assert_called_once_with(attendance)

    @patch("attendance.views.clock_in_out.GraceTime")
    @patch("attendance.views.clock_in_out.enable_late_come_early_out_tracking")
    @patch("attendance.views.clock_in_out.early_out_create")
    def test_early_out_uses_attendance_clock_time(
        self, early_out_create, tracking, grace_time
    ):
        tracking.return_value = {"tracking": True}
        grace_time.objects.filter.return_value.exists.return_value = False
        attendance = SimpleNamespace(attendance_clock_out=time(16, 0))

        clock_in_out_views.early_out(
            attendance=attendance,
            start_time=attendance_utils.strtime_seconds("09:00"),
            end_time=attendance_utils.strtime_seconds("17:00"),
            shift=None,
        )

        early_out_create.assert_called_once_with(attendance)


class PortalClockOutTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 4, 10, 17, 0, 0))
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_portal_clock_out_auto_validates_without_auto_approving_overtime(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _now_mock,
        clock_out_helper_mock,
        _reverse_geocode_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/clock-out/",
            {"employee_id": "1", "latitude": "14.6", "longitude": "121.0"},
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }

        employee = MagicMock()
        employee.id = 1
        employee.get_full_name.return_value = "Test Employee"
        employee_model.objects.get.return_value = employee

        open_activity_qs = MagicMock()
        open_activity = MagicMock()
        open_activity.location_verified = True
        open_activity.clock_out = datetime(2026, 4, 10, 17, 0, 0)
        open_activity_qs.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value = open_activity_qs

        attendance = MagicMock()
        attendance.attendance_date = date(2026, 4, 10)
        attendance.attendance_worked_hour = "08:00"
        attendance.attendance_validated = False
        clock_out_helper_mock.return_value = attendance

        response = public_clock_out(request)

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertTrue(payload["success"])
        clock_out_helper_mock.assert_called_once()
        self.assertTrue(clock_out_helper_mock.call_args.kwargs["auto_validate"])
        self.assertFalse(clock_out_helper_mock.call_args.kwargs["auto_approve_overtime"])

    @patch("attendance.views.portal._reverse_geocode", return_value="Clock Out Address")
    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 4, 10, 17, 0, 0))
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_portal_clock_out_updates_open_activity_record_locations(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _now_mock,
        clock_out_helper_mock,
        _reverse_geocode_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/clock-out/",
            {"employee_id": "1", "latitude": "14.6", "longitude": "121.0"},
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }

        employee = MagicMock()
        employee.id = 1
        employee.get_full_name.return_value = "Test Employee"
        employee_model.objects.get.return_value = employee

        open_activity_qs = MagicMock()
        open_activity = MagicMock()
        open_activity.location_verified = True
        open_activity.clock_out = datetime(2026, 4, 10, 17, 0, 0)
        open_activity.clock_in_gps_address = "Clock In Address"
        open_activity.clock_in_latitude = 14.5001
        open_activity.clock_in_longitude = 120.9001
        open_activity.clock_out_gps_address = None
        open_activity.clock_out_latitude = None
        open_activity.clock_out_longitude = None
        open_activity_qs.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value = open_activity_qs

        attendance = MagicMock()
        attendance.attendance_date = date(2026, 4, 10)
        attendance.attendance_worked_hour = "08:00"
        clock_out_helper_mock.return_value = attendance

        response = public_clock_out(request)

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertTrue(payload["success"])
        open_activity.refresh_from_db.assert_called_once()
        self.assertEqual(open_activity.clock_in_gps_address, "Clock In Address")
        self.assertEqual(open_activity.clock_out_gps_address, "Clock Out Address")
        self.assertEqual(open_activity.clock_out_latitude, 14.6)
        self.assertEqual(open_activity.clock_out_longitude, 121.0)

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 4, 10, 17, 0, 0))
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_portal_clock_out_invalid_coordinates_preserve_clock_in_location(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _now_mock,
        clock_out_helper_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/clock-out/",
            {"employee_id": "1", "latitude": "invalid", "longitude": "invalid"},
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }

        employee = MagicMock()
        employee.id = 1
        employee.get_full_name.return_value = "Test Employee"
        employee_model.objects.get.return_value = employee

        open_activity_qs = MagicMock()
        open_activity = MagicMock()
        open_activity.location_verified = True
        open_activity.clock_out = datetime(2026, 4, 10, 17, 0, 0)
        open_activity.clock_in_gps_address = "Clock In Address"
        open_activity.clock_in_latitude = 14.5001
        open_activity.clock_in_longitude = 120.9001
        open_activity.clock_out_gps_address = None
        open_activity.clock_out_latitude = None
        open_activity.clock_out_longitude = None
        open_activity_qs.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value = open_activity_qs

        attendance = MagicMock()
        attendance.attendance_date = date(2026, 4, 10)
        attendance.attendance_worked_hour = "08:00"
        clock_out_helper_mock.return_value = attendance

        response = public_clock_out(request)

        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertTrue(payload["success"])
        self.assertEqual(open_activity.clock_in_gps_address, "Clock In Address")
        self.assertEqual(open_activity.clock_in_latitude, 14.5001)
        self.assertEqual(open_activity.clock_in_longitude, 120.9001)
        self.assertIsNone(open_activity.clock_out_gps_address)
        self.assertTrue(open_activity.location_verified)


class PortalGeofenceCheckTests(SimpleTestCase):
    def _employee(self, geofences):
        return SimpleNamespace(id=1, pk=1, active_assigned_geofences=geofences)

    def _geofence(self, *, start=True, excluded_employees=None, radius=100):
        return SimpleNamespace(
            name="HQ",
            latitude=14.6000,
            longitude=121.0000,
            radius_in_meters=radius,
            start=start,
            excluded_employees=excluded_employees or [],
        )

    def test_geofence_check_allows_employee_with_no_assigned_geofences(self):
        employee = self._employee([])

        self.assertIsNone(_geofence_check(employee, None, "15.0", "122.0"))

    def test_geofence_check_allows_employee_with_only_inactive_geofences(self):
        employee = self._employee([self._geofence(start=False)])

        self.assertIsNone(_geofence_check(employee, None, "15.0", "122.0"))

    def test_geofence_check_allows_employee_excluded_from_active_geofence(self):
        employee = self._employee([])
        employee.active_assigned_geofences = [
            self._geofence(excluded_employees=[SimpleNamespace(id=1, pk=1)])
        ]

        self.assertIsNone(_geofence_check(employee, None, "15.0", "122.0"))

    def test_geofence_check_blocks_employee_outside_active_geofence(self):
        employee = self._employee([self._geofence()])

        result = _geofence_check(employee, None, "15.0", "122.0")

        self.assertIsNotNone(result)
        self.assertEqual(result["user_lat"], 15.0)
        self.assertEqual(result["geo_radius_meters"], 100)

    def test_geofence_check_allows_employee_inside_active_geofence(self):
        employee = self._employee([self._geofence()])

        self.assertIsNone(_geofence_check(employee, None, "14.6001", "121.0001"))


class PortalUnrestrictedGeofenceEndpointTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _verified_request(self, path, data):
        request = self.factory.post(path, data)
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": data["employee_id"],
            "verified_at": datetime.now().timestamp(),
        }
        return request

    def _employee(self):
        employee = MagicMock()
        employee.id = 1
        employee.pk = 1
        employee.active_assigned_geofences = []
        employee.employee_work_info = MagicMock()
        employee.employee_work_info.shift_id = object()
        employee.get_full_name.return_value = "Test Employee"
        return employee

    @patch("attendance.views.portal._maybe_auto_checkout_employee")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal.clock_in_attendance_and_activity")
    @patch("attendance.views.portal.shift_schedule_today", return_value=("08:00", 32400, 61200))
    @patch("attendance.views.portal.EmployeeShiftDay")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 9, 0, 0))
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_unrestricted_employee_can_clock_in_outside_geofence(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _now_mock,
        shift_day_model,
        _shift_schedule_mock,
        _clock_in_helper_mock,
        _reverse_mock,
        _auto_checkout_mock,
    ):
        request = self._verified_request(
            "/attendance/portal/clock-in/",
            {"employee_id": "1", "latitude": "15.0", "longitude": "122.0"},
        )
        employee_model.objects.get.return_value = self._employee()
        shift_day_model.objects.filter.return_value.first.return_value = object()
        activity = MagicMock()
        activity.location_verified = False
        activity_qs = MagicMock()
        activity_qs.exists.return_value = False
        activity_qs.order_by.return_value.last.return_value = activity
        attendance_activity_model.objects.filter.return_value = activity_qs

        response = public_clock_in(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(activity.clock_in_latitude, 15.0)
        self.assertEqual(activity.clock_in_longitude, 122.0)

    @patch("attendance.views.portal._maybe_auto_checkout_employee")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 17, 0, 0))
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_unrestricted_employee_can_clock_out_outside_geofence(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _now_mock,
        clock_out_helper_mock,
        _reverse_mock,
        _auto_checkout_mock,
    ):
        request = self._verified_request(
            "/attendance/portal/clock-out/",
            {"employee_id": "1", "latitude": "15.0", "longitude": "122.0"},
        )
        employee_model.objects.get.return_value = self._employee()
        open_activity = MagicMock()
        open_activity.activity_type = "work"
        open_activity.location_verified = True
        open_activity.clock_out = datetime(2026, 6, 5, 17, 0, 0)
        activity_qs = MagicMock()
        activity_qs.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value = activity_qs
        attendance = MagicMock()
        attendance.attendance_worked_hour = "08:00"
        clock_out_helper_mock.return_value = attendance

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(open_activity.clock_out_latitude, 15.0)
        self.assertEqual(open_activity.clock_out_longitude, 122.0)

    @patch("attendance.views.portal._maybe_auto_checkout_employee")
    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 0, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_unrestricted_employee_can_start_break_outside_geofence(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _reverse_mock,
        _worked_hours_mock,
        _atomic_mock,
        _auto_checkout_mock,
    ):
        request = self._verified_request(
            "/attendance/portal/activity-transition/",
            {
                "employee_id": "1",
                "activity_type": "break",
                "transition": "start",
                "latitude": "15.0",
                "longitude": "122.0",
            },
        )
        employee_model.objects.get.return_value = self._employee()
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "work"
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.exists.return_value = False
        attendance_activity_model.objects.filter.return_value.count.return_value = 0
        attendance_activity_model.objects.create.return_value = MagicMock()

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["activity_type"], "break")
        self.assertEqual(payload["transition"], "start")

    @patch("attendance.views.portal._maybe_auto_checkout_employee")
    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 13, 0, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_unrestricted_employee_can_end_lunch_outside_geofence(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _reverse_mock,
        _worked_hours_mock,
        _atomic_mock,
        _auto_checkout_mock,
    ):
        request = self._verified_request(
            "/attendance/portal/activity-transition/",
            {
                "employee_id": "1",
                "activity_type": "lunch",
                "transition": "end",
                "latitude": "15.0",
                "longitude": "122.0",
            },
        )
        employee_model.objects.get.return_value = self._employee()
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "lunch"
        open_activity.in_datetime = datetime(2026, 6, 5, 12, 0, 0)
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.exists.return_value = True
        attendance_activity_model.objects.filter.return_value.count.return_value = 0
        attendance_activity_model.objects.create.return_value = MagicMock()

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["activity_type"], "lunch")
        self.assertEqual(payload["transition"], "end")


class PortalAutoCheckoutTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _auto_checkout_context(self, schedule=None, activity_type="work"):
        employee = MagicMock()
        shift = object()
        shift_day = object()

        activity = SimpleNamespace(
            attendance_date=date(2026, 6, 5),
            clock_in_date=date(2026, 6, 5),
            clock_in=time(9, 0),
            in_datetime=datetime(2026, 6, 5, 9, 0, 0),
            shift_day=shift_day,
            activity_type=activity_type,
        )
        attendance = SimpleNamespace(
            shift_id=shift,
            attendance_day=shift_day,
        )
        schedule = schedule or SimpleNamespace(
            is_auto_punch_out_enabled=True,
            auto_punch_out_time=time(17, 30),
            start_time=time(9, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        return employee, activity, attendance, schedule

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.EmployeeShiftSchedule")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal._open_portal_activity")
    def test_auto_checkout_closes_open_work_activity_after_configured_time(
        self,
        open_activity_mock,
        attendance_model,
        schedule_model,
        clock_out_helper,
    ):
        employee, activity, attendance, schedule = self._auto_checkout_context()
        open_activity_mock.return_value = activity
        attendance_model.objects.filter.return_value.first.return_value = attendance
        schedule_model.objects.filter.return_value.first.return_value = schedule
        clock_out_helper.return_value = attendance

        result = _maybe_auto_checkout_employee(
            employee,
            current_time=datetime(2026, 6, 5, 18, 0, 0),
        )

        self.assertEqual(result, attendance)
        clock_out_helper.assert_called_once_with(
            employee=employee,
            date_today=date(2026, 6, 5),
            now="17:30",
            out_datetime=datetime(2026, 6, 5, 17, 30, 0),
            auto_validate=True,
            auto_approve_overtime=False,
        )

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.EmployeeShiftSchedule")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal._open_portal_activity")
    def test_auto_checkout_does_nothing_before_configured_time(
        self,
        open_activity_mock,
        attendance_model,
        schedule_model,
        clock_out_helper,
    ):
        employee, activity, attendance, schedule = self._auto_checkout_context()
        open_activity_mock.return_value = activity
        attendance_model.objects.filter.return_value.first.return_value = attendance
        schedule_model.objects.filter.return_value.first.return_value = schedule

        result = _maybe_auto_checkout_employee(
            employee,
            current_time=datetime(2026, 6, 5, 17, 0, 0),
        )

        self.assertIsNone(result)
        clock_out_helper.assert_not_called()

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.EmployeeShiftSchedule")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal._open_portal_activity")
    def test_auto_checkout_requires_enabled_schedule_and_time(
        self,
        open_activity_mock,
        attendance_model,
        schedule_model,
        clock_out_helper,
    ):
        schedules = [
            SimpleNamespace(
                is_auto_punch_out_enabled=False,
                auto_punch_out_time=time(17, 30),
                start_time=time(9, 0),
                end_time=time(17, 0),
                is_night_shift=False,
            ),
            SimpleNamespace(
                is_auto_punch_out_enabled=True,
                auto_punch_out_time=None,
                start_time=time(9, 0),
                end_time=time(17, 0),
                is_night_shift=False,
            ),
        ]

        for schedule in schedules:
            with self.subTest(schedule=schedule):
                employee, activity, attendance, _schedule = self._auto_checkout_context(
                    schedule=schedule
                )
                open_activity_mock.return_value = activity
                attendance_model.objects.filter.return_value.first.return_value = attendance
                schedule_model.objects.filter.return_value.first.return_value = schedule

                result = _maybe_auto_checkout_employee(
                    employee,
                    current_time=datetime(2026, 6, 5, 18, 0, 0),
                )

                self.assertIsNone(result)

        clock_out_helper.assert_not_called()

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.EmployeeShiftSchedule")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal._open_portal_activity")
    def test_auto_checkout_uses_next_day_for_night_shift_checkout_time(
        self,
        open_activity_mock,
        attendance_model,
        schedule_model,
        clock_out_helper,
    ):
        schedule = SimpleNamespace(
            is_auto_punch_out_enabled=True,
            auto_punch_out_time=time(6, 30),
            start_time=time(22, 0),
            end_time=time(6, 0),
            is_night_shift=True,
        )
        employee, activity, attendance, _schedule = self._auto_checkout_context(
            schedule=schedule
        )
        activity.clock_in = time(22, 0)
        activity.in_datetime = datetime(2026, 6, 5, 22, 0, 0)
        open_activity_mock.return_value = activity
        attendance_model.objects.filter.return_value.first.return_value = attendance
        schedule_model.objects.filter.return_value.first.return_value = schedule
        clock_out_helper.return_value = attendance

        result = _maybe_auto_checkout_employee(
            employee,
            current_time=datetime(2026, 6, 6, 7, 0, 0),
        )

        self.assertEqual(result, attendance)
        clock_out_helper.assert_called_once_with(
            employee=employee,
            date_today=date(2026, 6, 6),
            now="06:30",
            out_datetime=datetime(2026, 6, 6, 6, 30, 0),
            auto_validate=True,
            auto_approve_overtime=False,
        )

    @patch("attendance.views.portal.clock_out_attendance_and_activity")
    @patch("attendance.views.portal.EmployeeShiftSchedule")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal._open_portal_activity")
    def test_auto_checkout_skips_non_work_activity(
        self,
        open_activity_mock,
        attendance_model,
        schedule_model,
        clock_out_helper,
    ):
        employee, activity, attendance, schedule = self._auto_checkout_context(
            activity_type="break"
        )
        open_activity_mock.return_value = activity
        attendance_model.objects.filter.return_value.first.return_value = attendance
        schedule_model.objects.filter.return_value.first.return_value = schedule

        result = _maybe_auto_checkout_employee(
            employee,
            current_time=datetime(2026, 6, 5, 18, 0, 0),
        )

        self.assertIsNone(result)
        clock_out_helper.assert_not_called()

    @patch("attendance.views.portal._portal_worked_duration_metadata", return_value={"worked_total_seconds": 30600, "worked_total_time": "08:30"})
    @patch("attendance.views.portal._portal_activity_duration_metadata", return_value={"break_total_seconds": 0, "break_total_time": "00:00", "lunch_total_seconds": 0, "lunch_total_time": "00:00"})
    @patch("attendance.views.portal._portal_break_metadata", return_value={"breaks_taken": 0, "breaks_allowed": 2, "break_minutes_allowed": 15, "lunch_minutes_allowed": 60, "break_limit_reached": False})
    @patch("attendance.views.portal._lunch_taken", return_value=False)
    @patch("attendance.views.portal._latest_portal_activity")
    @patch("attendance.views.portal._open_portal_activity", return_value=None)
    @patch("attendance.views.portal._maybe_auto_checkout_employee")
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_reflects_auto_closed_attendance(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        auto_checkout_mock,
        _open_activity_mock,
        latest_activity_mock,
        _lunch_taken_mock,
        _break_metadata_mock,
        _activity_duration_mock,
        _worked_duration_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "1234567"},
        )
        attach_session(request)

        employee = MagicMock()
        employee.id = 1
        employee.employee_no = "1234567"
        employee.employee_work_info = SimpleNamespace(branch_id=None)
        employee.get_full_name.return_value = "Test Employee"
        employee.get_avatar.return_value = "/avatar.png"
        employee_model.objects.filter.return_value = [employee]

        latest_activity_mock.return_value = SimpleNamespace(
            attendance_date=date(2026, 6, 5),
            clock_in_date=date(2026, 6, 5),
            clock_in=time(9, 0),
            in_datetime=datetime(2026, 6, 5, 9, 0, 0),
            clock_out_date=date(2026, 6, 5),
            clock_out=time(17, 30),
            out_datetime=datetime(2026, 6, 5, 17, 30, 0),
        )
        attendance_model.objects.filter.return_value.first.return_value = SimpleNamespace(
            attendance_clock_in_date=date(2026, 6, 5),
            attendance_clock_in=time(9, 0),
            attendance_clock_out_date=date(2026, 6, 5),
            attendance_clock_out=time(17, 30),
        )

        response = employee_lookup(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertFalse(payload["results"][0]["is_clocked_in"])
        self.assertEqual(payload["results"][0]["clock_out_datetime"], "2026-06-05T17:30:00")
        auto_checkout_mock.assert_called_once_with(employee)


class PortalActivityDurationTests(SimpleTestCase):
    def _activity(self, start, end=None, clock_out=None):
        return SimpleNamespace(
            in_datetime=start,
            out_datetime=end,
            clock_in_date=None,
            clock_in=None,
            clock_out_date=None,
            clock_out=clock_out,
        )

    @patch("attendance.views.portal.AttendanceActivity")
    def test_activity_total_sums_multiple_break_rows(self, attendance_activity_model):
        employee = object()
        attendance_date = date(2026, 6, 5)
        attendance_activity_model.objects.filter.return_value = [
            self._activity(
                datetime(2026, 6, 5, 10, 0, 0),
                datetime(2026, 6, 5, 10, 15, 0),
            ),
            self._activity(
                datetime(2026, 6, 5, 15, 0, 0),
                datetime(2026, 6, 5, 15, 10, 0),
            ),
        ]

        total_seconds = _portal_activity_total_seconds(
            employee,
            attendance_date,
            "break",
            datetime(2026, 6, 5, 16, 0, 0),
        )

        self.assertEqual(total_seconds, 1500)
        attendance_activity_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date=attendance_date,
            activity_type="break",
        )

    @patch("attendance.views.portal.AttendanceActivity")
    def test_activity_total_sums_lunch_rows(self, attendance_activity_model):
        employee = object()
        attendance_date = date(2026, 6, 5)
        attendance_activity_model.objects.filter.return_value = [
            self._activity(
                datetime(2026, 6, 5, 12, 0, 0),
                datetime(2026, 6, 5, 12, 45, 0),
            ),
        ]

        total_seconds = _portal_activity_total_seconds(
            employee,
            attendance_date,
            "lunch",
            datetime(2026, 6, 5, 16, 0, 0),
        )

        self.assertEqual(total_seconds, 2700)
        attendance_activity_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date=attendance_date,
            activity_type="lunch",
        )

    @patch("attendance.views.portal.AttendanceActivity")
    def test_activity_total_ignores_missing_or_invalid_times(
        self,
        attendance_activity_model,
    ):
        attendance_activity_model.objects.filter.return_value = [
            self._activity(None, datetime(2026, 6, 5, 10, 15, 0)),
            self._activity("bad-start", "bad-end", clock_out="bad-end"),
        ]

        total_seconds = _portal_activity_total_seconds(
            object(),
            date(2026, 6, 5),
            "break",
            datetime(2026, 6, 5, 16, 0, 0),
        )

        self.assertEqual(total_seconds, 0)

    @patch("attendance.views.portal.AttendanceActivity")
    def test_activity_total_includes_open_current_activity(
        self,
        attendance_activity_model,
    ):
        attendance_activity_model.objects.filter.return_value = [
            self._activity(datetime(2026, 6, 5, 14, 0, 0)),
        ]

        total_seconds = _portal_activity_total_seconds(
            object(),
            date(2026, 6, 5),
            "break",
            datetime(2026, 6, 5, 14, 20, 0),
        )

        self.assertEqual(total_seconds, 1200)

    @patch("attendance.views.portal.Attendance")
    def test_worked_duration_uses_attendance_clock_times(self, attendance_model):
        employee = object()
        attendance_date = date(2026, 6, 5)
        attendance_model.objects.filter.return_value.first.return_value = SimpleNamespace(
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=datetime(2026, 6, 5, 9, 0, 0).time(),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=datetime(2026, 6, 5, 17, 0, 0).time(),
        )

        metadata = _portal_worked_duration_metadata(
            employee,
            attendance_date,
            datetime(2026, 6, 5, 18, 0, 0),
        )

        self.assertEqual(metadata["worked_total_seconds"], 28800)
        self.assertEqual(metadata["worked_total_time"], "08:00")
        attendance_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date=attendance_date,
        )

    @patch("attendance.views.portal.Attendance")
    def test_worked_duration_does_not_subtract_lunch_or_break(self, attendance_model):
        attendance_date = date(2026, 6, 5)
        attendance_model.objects.filter.return_value.first.return_value = SimpleNamespace(
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=datetime(2026, 6, 5, 9, 0, 0).time(),
            attendance_clock_out_date=None,
            attendance_clock_out=None,
        )

        metadata = _portal_worked_duration_metadata(
            object(),
            attendance_date,
            datetime(2026, 6, 5, 14, 30, 0),
        )

        self.assertEqual(metadata["worked_total_seconds"], 19800)
        self.assertEqual(metadata["worked_total_time"], "05:30")


class PortalEmployeeLookupTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _employee(self):
        employee = MagicMock()
        employee.id = 1
        employee.employee_no = "0000001"
        employee.is_active = True
        employee.employee_work_info = None
        employee.get_full_name.return_value = "Test Employee"
        employee.get_avatar.return_value = "/avatar.png"
        return employee

    def _activity_qs(self, activity=None, activities=None):
        if activities is None:
            activities = [activity] if activity else []
        first_activity = activity if activity is not None else (activities[0] if activities else None)
        qs = MagicMock()
        qs.order_by.return_value.first.return_value = first_activity
        qs.exists.return_value = bool(activities)
        qs.count.return_value = len(activities)
        qs.__iter__.return_value = iter(activities)
        return qs

    def _attendance(self, clock_in, clock_out=None):
        return SimpleNamespace(
            attendance_date=date(2026, 6, 5),
            attendance_clock_in_date=date(2026, 6, 5),
            attendance_clock_in=clock_in.time(),
            attendance_clock_out_date=date(2026, 6, 5) if clock_out else None,
            attendance_clock_out=clock_out.time() if clock_out else None,
        )

    def _activity_filter(self, open_activity=None, latest_activity=None, work_activities=None, break_activities=None, lunch_activities=None):
        work_activities = work_activities or []
        break_activities = break_activities or []
        lunch_activities = lunch_activities or []

        def filter_side_effect(*args, **kwargs):
            if kwargs.get("clock_out__isnull") is True:
                return self._activity_qs(open_activity)
            activity_type = kwargs.get("activity_type")
            if activity_type == "work":
                return self._activity_qs(activities=work_activities)
            if activity_type == "break":
                return self._activity_qs(activities=break_activities)
            if activity_type == "lunch":
                return self._activity_qs(activities=lunch_activities)
            return self._activity_qs(latest_activity)

        return filter_side_effect

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_returns_open_clock_in_time(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001"},
        )
        attach_session(request)
        employee = self._employee()
        employee_model.objects.filter.return_value = [employee]

        open_activity = SimpleNamespace(
            in_datetime=datetime(2026, 6, 5, 9, 45, 0),
            out_datetime=None,
            clock_in_date=None,
            clock_in=None,
            clock_out_date=None,
            clock_out=None,
            attendance_date=date(2026, 6, 5),
            activity_type="work",
        )
        attendance_model.objects.filter.return_value.first.return_value = self._attendance(
            datetime(2026, 6, 5, 9, 45, 0)
        )
        attendance_activity_model.objects.filter.side_effect = self._activity_filter(
            open_activity=open_activity,
            work_activities=[open_activity],
        )

        response = employee_lookup(request)
        payload = json.loads(response.content)
        result = payload["results"][0]

        self.assertTrue(result["is_clocked_in"])
        self.assertEqual(result["clock_in_datetime"], "2026-06-05T09:45:00")
        self.assertIsNone(result["clock_out_datetime"])
        self.assertEqual(result["breaks_taken"], 0)
        self.assertEqual(result["breaks_allowed"], 2)
        self.assertEqual(result["break_minutes_allowed"], 15)
        self.assertEqual(result["lunch_minutes_allowed"], 60)
        self.assertFalse(result["break_limit_reached"])
        self.assertEqual(result["break_total_seconds"], 0)
        self.assertEqual(result["break_total_time"], "00:00")
        self.assertEqual(result["lunch_total_seconds"], 0)
        self.assertEqual(result["lunch_total_time"], "00:00")

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_returns_latest_closed_clock_times(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001"},
        )
        attach_session(request)
        employee = self._employee()
        employee_model.objects.filter.return_value = [employee]

        closed_activity = SimpleNamespace(
            in_datetime=datetime(2026, 6, 5, 9, 45, 0),
            out_datetime=datetime(2026, 6, 5, 17, 30, 0),
            clock_in_date=None,
            clock_in=None,
            clock_out_date=None,
            clock_out=None,
            attendance_date=date(2026, 6, 5),
            activity_type="work",
        )
        attendance_model.objects.filter.return_value.first.return_value = self._attendance(
            datetime(2026, 6, 5, 9, 45, 0),
            datetime(2026, 6, 5, 17, 30, 0),
        )
        attendance_activity_model.objects.filter.side_effect = self._activity_filter(
            latest_activity=closed_activity,
            work_activities=[closed_activity],
        )

        response = employee_lookup(request)
        payload = json.loads(response.content)
        result = payload["results"][0]

        self.assertFalse(result["is_clocked_in"])
        self.assertEqual(result["clock_in_datetime"], "2026-06-05T09:45:00")
        self.assertEqual(result["clock_out_datetime"], "2026-06-05T17:30:00")
        self.assertEqual(result["break_total_time"], "00:00")
        self.assertEqual(result["lunch_total_time"], "00:00")

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    @patch("attendance.views.portal.timezone.now", return_value=datetime(2026, 6, 5, 12, 30, 0))
    def test_employee_lookup_uses_attendance_clock_in_while_on_lunch(
        self,
        _now_mock,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001"},
        )
        attach_session(request)
        employee = self._employee()
        employee_model.objects.filter.return_value = [employee]
        attendance_model.objects.filter.return_value.first.return_value = self._attendance(
            datetime(2026, 6, 5, 9, 0, 0)
        )

        work_activity = SimpleNamespace(
            in_datetime=datetime(2026, 6, 5, 9, 0, 0),
            out_datetime=datetime(2026, 6, 5, 12, 0, 0),
            clock_in_date=None,
            clock_in=None,
            clock_out_date=None,
            clock_out=None,
            attendance_date=date(2026, 6, 5),
            activity_type="work",
        )
        lunch_activity = SimpleNamespace(
            in_datetime=datetime(2026, 6, 5, 12, 0, 0),
            out_datetime=None,
            clock_in_date=None,
            clock_in=None,
            clock_out_date=None,
            clock_out=None,
            attendance_date=date(2026, 6, 5),
            activity_type="lunch",
        )
        attendance_activity_model.objects.filter.side_effect = self._activity_filter(
            open_activity=lunch_activity,
            work_activities=[work_activity],
            lunch_activities=[lunch_activity],
        )

        response = employee_lookup(request)
        payload = json.loads(response.content)
        result = payload["results"][0]

        self.assertTrue(result["is_clocked_in"])
        self.assertEqual(result["active_activity_type"], "lunch")
        self.assertEqual(result["clock_in_datetime"], "2026-06-05T09:00:00")
        self.assertEqual(result["worked_total_seconds"], 12600)

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_with_suffix(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        from django.db.models import Q
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001-2"},
        )
        attach_session(request)
        employee = self._employee()
        employee.employee_no = "0000001-2"
        employee_model.objects.filter.return_value = [employee]
        attendance_model.objects.filter.return_value.first.return_value = None
        attendance_activity_model.objects.filter.side_effect = self._activity_filter()

        response = employee_lookup(request)
        payload = json.loads(response.content)
        
        self.assertTrue(payload["success"])
        self.assertEqual(len(payload["results"]), 1)
        self.assertEqual(payload["results"][0]["employee_no"], "0000001-2")
        
        # Verify the Django Q filter constructed for exact match
        args, kwargs = employee_model.objects.filter.call_args
        self.assertTrue(any(isinstance(arg, Q) for arg in args))
        q_obj = next(arg for arg in args if isinstance(arg, Q))
        self.assertEqual(len(q_obj.children), 1)
        self.assertEqual(q_obj.children[0], ("employee_no__exact", "0000001-2"))

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_without_suffix_matches_with_or_without_suffix(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        from django.db.models import Q
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001"},
        )
        attach_session(request)
        employee = self._employee()
        employee.employee_no = "0000001-2"
        employee_model.objects.filter.return_value = [employee]
        attendance_model.objects.filter.return_value.first.return_value = None
        attendance_activity_model.objects.filter.side_effect = self._activity_filter()

        response = employee_lookup(request)
        payload = json.loads(response.content)
        
        self.assertTrue(payload["success"])
        self.assertEqual(len(payload["results"]), 1)
        self.assertEqual(payload["results"][0]["employee_no"], "0000001-2")
        
        # Verify the Django Q filter constructed for exact OR prefix with hyphen
        args, kwargs = employee_model.objects.filter.call_args
        self.assertTrue(any(isinstance(arg, Q) for arg in args))
        q_obj = next(arg for arg in args if isinstance(arg, Q))
        self.assertEqual(q_obj.connector, Q.OR)
        self.assertEqual(len(q_obj.children), 2)
        filters = dict(q_obj.children)
        self.assertEqual(filters["employee_no__exact"], "0000001")
        self.assertEqual(filters["employee_no__startswith"], "0000001-")

    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_employee_lookup_geofence_preview_uses_only_enforced_geofences(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
    ):
        request = self.factory.post(
            "/attendance/portal/employee-lookup/",
            {"query": "0000001"},
        )
        attach_session(request)
        employee = self._employee()
        employee.pk = 1
        employee.active_assigned_geofences = [
            SimpleNamespace(
                name="HQ",
                latitude=14.6,
                longitude=121.0,
                radius_in_meters=100,
                start=True,
                excluded_employees=[SimpleNamespace(id=1, pk=1)],
            )
        ]
        employee_model.objects.filter.return_value = [employee]
        attendance_model.objects.filter.return_value.first.return_value = None
        attendance_activity_model.objects.filter.side_effect = self._activity_filter()

        response = employee_lookup(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["results"][0]["geo_fence"], [])


class PortalAttendanceHistoryTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _verified_request(self, employee_id="1"):
        request = self.factory.post(
            "/attendance/portal/attendance-history/",
            {"employee_id": employee_id},
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": employee_id,
            "verified_at": datetime.now().timestamp(),
        }
        return request

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_requires_verified_pin(self, _ip_allowed_mock):
        request = self.factory.post(
            "/attendance/portal/attendance-history/",
            {"employee_id": "1"},
        )
        attach_session(request)

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("PIN verification required", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_rejects_mismatched_pin(self, _ip_allowed_mock):
        request = self._verified_request(employee_id="1")
        request.session["portal_pin_verification"]["employee_id"] = "2"

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("does not match", payload["message"])

    @patch("attendance.views.portal.timezone.localdate", return_value=date(2026, 6, 5))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_allows_old_verified_pin(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        _localdate_mock,
    ):
        request = self._verified_request(employee_id="1")
        request.session["portal_pin_verification"]["verified_at"] = (
            datetime.now() - timedelta(minutes=5)
        ).timestamp()
        employee_model.objects.get.return_value = MagicMock()
        attendance_model.objects.filter.return_value.order_by.return_value = []

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["rows"], [])

    @patch("attendance.views.portal.timezone.localdate", return_value=date(2026, 6, 5))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_returns_default_last_15_day_rows(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        _localdate_mock,
    ):
        request = self._verified_request(employee_id="1")
        employee = MagicMock()
        employee_model.objects.get.return_value = employee
        attendance_rows = [
            SimpleNamespace(
                attendance_date=date(2026, 6, 5),
                attendance_clock_in=datetime(2026, 6, 5, 9, 45).time(),
                attendance_clock_out=datetime(2026, 6, 5, 17, 30).time(),
                attendance_worked_hour="07:45",
            ),
            SimpleNamespace(
                attendance_date=date(2026, 6, 4),
                attendance_clock_in=datetime(2026, 6, 4, 9, 0).time(),
                attendance_clock_out=None,
                attendance_worked_hour=None,
            ),
        ]
        attendance_model.objects.filter.return_value.order_by.return_value = attendance_rows

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        attendance_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date__gte=date(2026, 5, 22),
            attendance_date__lte=date(2026, 6, 5),
        )
        self.assertEqual(payload["start_date"], "2026-05-22")
        self.assertEqual(payload["end_date"], "2026-06-05")
        self.assertEqual(
            payload["rows"][0],
            {
                "attendance_date": "Jun 05, 2026",
                "clock_in_time": "09:45 AM",
                "clock_out_time": "05:30 PM",
                "worked_hours": "07:45",
                "break_time": "00:00",
                "lunch_time": "00:00",
            },
        )
        self.assertEqual(payload["rows"][1]["clock_out_time"], "--:--")
        self.assertEqual(payload["rows"][1]["worked_hours"], "00:00")
        self.assertEqual(payload["rows"][1]["break_time"], "00:00")
        self.assertEqual(payload["rows"][1]["lunch_time"], "00:00")

    @patch("attendance.views.portal.timezone.localdate", return_value=date(2026, 6, 5))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_uses_custom_date_range(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        _localdate_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/attendance-history/",
            {
                "employee_id": "1",
                "start_date": "2026-06-01",
                "end_date": "2026-06-03",
            },
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }
        employee = MagicMock()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.order_by.return_value = []

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        attendance_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date__gte=date(2026, 6, 1),
            attendance_date__lte=date(2026, 6, 3),
        )
        self.assertEqual(payload["start_date"], "2026-06-01")
        self.assertEqual(payload["end_date"], "2026-06-03")

    @patch("attendance.views.portal.timezone.localdate", return_value=date(2026, 6, 5))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_invalid_dates_fall_back_to_default(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        _localdate_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/attendance-history/",
            {
                "employee_id": "1",
                "start_date": "not-a-date",
                "end_date": "also-bad",
            },
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }
        employee = MagicMock()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.order_by.return_value = []

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        attendance_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date__gte=date(2026, 5, 22),
            attendance_date__lte=date(2026, 6, 5),
        )
        self.assertEqual(payload["start_date"], "2026-05-22")
        self.assertEqual(payload["end_date"], "2026-06-05")

    @patch("attendance.views.portal.timezone.localdate", return_value=date(2026, 6, 5))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_attendance_history_reversed_dates_are_swapped(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_model,
        _localdate_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/attendance-history/",
            {
                "employee_id": "1",
                "start_date": "2026-06-05",
                "end_date": "2026-06-01",
            },
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }
        employee = MagicMock()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.order_by.return_value = []

        response = attendance_history(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        attendance_model.objects.filter.assert_called_once_with(
            employee_id=employee,
            attendance_date__gte=date(2026, 6, 1),
            attendance_date__lte=date(2026, 6, 5),
        )
        self.assertEqual(payload["start_date"], "2026-06-01")
        self.assertEqual(payload["end_date"], "2026-06-05")


class PortalForgotPinResetTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _render_response(self, request, template, context):
        response = HttpResponse("rendered")
        response.template_name = template
        response.context_data = context
        return response

    def _employee(self, pin="123456"):
        work_info = MagicMock()
        work_info.pin = pin
        work_info.email = "work@example.com"

        employee = MagicMock()
        employee.pk = 1
        employee.id = 1
        employee.email = "personal@example.com"
        employee.employee_work_info = work_info
        employee.get_full_name.return_value = "Test Employee"
        return employee

    @patch("attendance.views.portal._send_portal_pin_reset_email", return_value=True)
    @patch("attendance.views.portal._get_portal_pin_reset_employees")
    def test_forgot_pin_accepts_personal_email(
        self,
        employees_mock,
        send_mock,
    ):
        employee = self._employee()
        employees_mock.return_value = [employee]
        request = self.factory.post(
            "/attendance/portal/forgot-pin/",
            {"email": "personal@example.com"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = forgot_pin(request)

        self.assertEqual(response.status_code, 200)
        employees_mock.assert_called_once_with("personal@example.com")
        send_mock.assert_called_once_with(request, employee)
        self.assertTrue(response.context_data["sent"])
        self.assertIn("active employee", response.context_data["generic_message"])

    @patch("attendance.views.portal._send_portal_pin_reset_email", return_value=True)
    @patch("attendance.views.portal._get_portal_pin_reset_employees")
    def test_forgot_pin_accepts_work_email(
        self,
        employees_mock,
        send_mock,
    ):
        employee = self._employee()
        employees_mock.return_value = [employee]
        request = self.factory.post(
            "/attendance/portal/forgot-pin/",
            {"email": "work@example.com"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = forgot_pin(request)

        self.assertEqual(response.status_code, 200)
        employees_mock.assert_called_once_with("work@example.com")
        send_mock.assert_called_once_with(request, employee)
        self.assertTrue(response.context_data["sent"])

    @patch("attendance.views.portal._send_portal_pin_reset_email", return_value=True)
    @patch("attendance.views.portal._get_portal_pin_reset_employees", return_value=[])
    def test_forgot_pin_returns_generic_success_when_no_employee_matches(
        self,
        employees_mock,
        send_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/forgot-pin/",
            {"email": "missing@example.com"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = forgot_pin(request)

        self.assertEqual(response.status_code, 200)
        employees_mock.assert_called_once_with("missing@example.com")
        send_mock.assert_not_called()
        self.assertTrue(response.context_data["sent"])
        self.assertIn("active employee", response.context_data["generic_message"])

    @patch("attendance.views.portal.render_to_string", return_value="<p>Reset</p>")
    @patch("attendance.views.portal.EmailMessage")
    @patch("attendance.views.portal.ConfiguredEmailBackend")
    def test_pin_reset_email_uses_configured_from_email(
        self,
        backend_class,
        email_message_class,
        _render_mock,
    ):
        employee = self._employee()
        request = self.factory.get(
            "/attendance/portal/forgot-pin/",
            HTTP_HOST="testserver",
        )
        backend = MagicMock()
        backend.dynamic_from_email_with_display_name = "MDC <no-reply@example.com>"
        backend_class.return_value = backend
        email_message = MagicMock()
        email_message_class.return_value = email_message

        sent = _send_portal_pin_reset_email(request, employee)

        self.assertTrue(sent)
        email_message_class.assert_called_once()
        self.assertEqual(
            email_message_class.call_args.kwargs["from_email"],
            "MDC <no-reply@example.com>",
        )
        self.assertEqual(email_message_class.call_args.kwargs["to"], ["work@example.com"])
        email_message.send.assert_called_once()

    def test_reset_token_loads_successfully(self):
        employee = self._employee()
        token = _make_portal_pin_reset_token(employee)

        with patch("attendance.views.portal.Employee.objects.get", return_value=employee):
            token_employee, token_error = _get_portal_pin_reset_employee(token)

        self.assertEqual(token_employee, employee)
        self.assertEqual(token_error, "")

    def test_reset_rejects_malformed_token(self):
        token_employee, token_error = _get_portal_pin_reset_employee("not-a-token")

        self.assertIsNone(token_employee)
        self.assertIn("invalid", token_error)

    @patch("attendance.views.portal.signing.loads", side_effect=SignatureExpired("old"))
    def test_reset_rejects_expired_token(self, _loads_mock):
        token_employee, token_error = _get_portal_pin_reset_employee("expired-token")

        self.assertIsNone(token_employee)
        self.assertIn("expired", token_error)

    def test_reset_rejects_invalidated_token_after_pin_changes(self):
        employee = self._employee(pin="123456")
        token = _make_portal_pin_reset_token(employee)
        employee.employee_work_info.pin = "654321"

        with patch("attendance.views.portal.Employee.objects.get", return_value=employee):
            token_employee, token_error = _get_portal_pin_reset_employee(token)

        self.assertIsNone(token_employee)
        self.assertIn("no longer valid", token_error)

    @patch("attendance.views.portal.Employee.objects.get")
    def test_reset_pin_requires_matching_six_digit_values(self, employee_get_mock):
        employee = self._employee()
        employee_get_mock.return_value = employee
        token = _make_portal_pin_reset_token(employee)
        request = self.factory.post(
            f"/attendance/portal/reset-pin/{token}/",
            {"new_pin": "123456", "confirm_pin": "654321"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = reset_pin(request, token)

        self.assertEqual(response.status_code, 200)
        self.assertIn("PIN values do not match", response.context_data["form"].errors.as_text())
        employee.employee_work_info.save.assert_not_called()

    @patch("attendance.views.portal.Employee.objects.get")
    def test_reset_pin_rejects_non_numeric_pin(self, employee_get_mock):
        employee = self._employee()
        employee_get_mock.return_value = employee
        token = _make_portal_pin_reset_token(employee)
        request = self.factory.post(
            f"/attendance/portal/reset-pin/{token}/",
            {"new_pin": "12AB56", "confirm_pin": "12AB56"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = reset_pin(request, token)

        self.assertEqual(response.status_code, 200)
        self.assertIn("PIN must be exactly 6 digits", response.context_data["form"].errors.as_text())
        employee.employee_work_info.save.assert_not_called()

    @patch("attendance.views.portal.Employee.objects.get")
    def test_reset_pin_updates_employee_work_info_pin(self, employee_get_mock):
        employee = self._employee()
        employee_get_mock.return_value = employee
        token = _make_portal_pin_reset_token(employee)
        request = self.factory.post(
            f"/attendance/portal/reset-pin/{token}/",
            {"new_pin": "654321", "confirm_pin": "654321"},
        )
        attach_session(request)

        with patch("attendance.views.portal.render", side_effect=self._render_response):
            response = reset_pin(request, token)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(employee.employee_work_info.pin, "654321")
        employee.employee_work_info.save.assert_called_once_with(update_fields=["pin"])
        self.assertTrue(response.context_data["reset_success"])


class PortalPinVerificationTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_verify_pin_success_sets_session(
        self,
        _ip_allowed_mock,
        employee_model,
    ):
        request = self.factory.post(
            "/attendance/portal/verify-pin/",
            {"employee_id": "1", "pin": "123456"},
        )
        attach_session(request)

        employee = MagicMock()
        employee.id = 1
        employee.employee_work_info = MagicMock(pin="123456")
        employee_model.objects.get.return_value = employee

        response = verify_pin(request)
        payload = json.loads(response.content)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(payload["success"])
        self.assertEqual(payload["attempts_remaining"], 3)
        self.assertEqual(
            request.session["portal_pin_verification"]["employee_id"],
            "1",
        )

    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_verify_pin_wrong_pin_tracks_attempts_and_resets_on_third(
        self,
        _ip_allowed_mock,
        employee_model,
    ):
        request = self.factory.post(
            "/attendance/portal/verify-pin/",
            {"employee_id": "1", "pin": "000000"},
        )
        attach_session(request)

        employee = MagicMock()
        employee.id = 1
        employee.employee_work_info = MagicMock(pin="123456")
        employee_model.objects.get.return_value = employee

        first = json.loads(verify_pin(request).content)
        second = json.loads(verify_pin(request).content)
        third = json.loads(verify_pin(request).content)

        self.assertFalse(first["success"])
        self.assertEqual(first["attempts_remaining"], 2)
        self.assertFalse(second["success"])
        self.assertEqual(second["attempts_remaining"], 1)
        self.assertFalse(third["success"])
        self.assertEqual(third["attempts_remaining"], 0)
        self.assertTrue(third["reset_required"])
        self.assertNotIn("1", request.session.get("portal_pin_attempts", {}))

    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_verify_pin_rejects_non_six_digit_pin(
        self,
        _ip_allowed_mock,
        employee_model,
    ):
        request = self.factory.post(
            "/attendance/portal/verify-pin/",
            {"employee_id": "1", "pin": "12345"},
        )
        attach_session(request)

        response = verify_pin(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertEqual(payload["attempts_remaining"], 3)
        self.assertIn("exactly 6 digits", payload["message"])
        employee_model.objects.get.assert_not_called()

    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_verify_pin_rejects_missing_stored_pin(
        self,
        _ip_allowed_mock,
        employee_model,
    ):
        request = self.factory.post(
            "/attendance/portal/verify-pin/",
            {"employee_id": "1", "pin": "123456"},
        )
        attach_session(request)

        employee = MagicMock()
        employee.id = 1
        employee.employee_work_info = MagicMock(pin=None)
        employee_model.objects.get.return_value = employee

        response = verify_pin(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("not configured", payload["message"])

    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_verify_pin_rejects_invalid_stored_pin_format(
        self,
        _ip_allowed_mock,
        employee_model,
    ):
        request = self.factory.post(
            "/attendance/portal/verify-pin/",
            {"employee_id": "1", "pin": "123456"},
        )
        attach_session(request)

        employee = MagicMock()
        employee.id = 1
        employee.employee_work_info = MagicMock(pin="12AB56")
        employee_model.objects.get.return_value = employee

        response = verify_pin(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("invalid", payload["message"])


class PortalActivityTransitionTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _verified_request(self, data=None):
        payload = {
            "employee_id": "1",
            "activity_type": "break",
            "transition": "start",
            "latitude": "14.6001",
            "longitude": "121.0001",
        }
        if data:
            payload.update(data)
        request = self.factory.post("/attendance/portal/activity-transition/", payload)
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": payload["employee_id"],
            "verified_at": datetime.now().timestamp(),
        }
        return request

    def _employee(self):
        employee = MagicMock()
        employee.id = 1
        employee.get_full_name.return_value = "Test Employee"
        employee.employee_work_info = MagicMock()
        return employee

    def test_break_policy_falls_back_to_defaults_without_settings(self):
        employee = self._employee()

        with patch(
            "attendance.views.portal.AttendanceGeneralSetting.objects.filter",
            side_effect=Exception("database unavailable"),
        ):
            policy = _portal_break_policy(employee)

        self.assertEqual(policy["breaks_allowed"], 2)
        self.assertEqual(policy["break_minutes_allowed"], 15)
        self.assertEqual(policy["lunch_minutes_allowed"], 60)

    def test_break_policy_prefers_company_setting(self):
        employee = self._employee()
        company = object()
        employee.employee_work_info.company_id = company
        setting = SimpleNamespace(
            portal_break_limit=3,
            portal_break_minutes=20,
            portal_lunch_minutes=45,
        )
        qs = MagicMock()
        qs.first.return_value = setting

        with patch(
            "attendance.views.portal.AttendanceGeneralSetting.objects.filter",
            return_value=qs,
        ) as filter_mock:
            policy = _portal_break_policy(employee)

        filter_mock.assert_called_once_with(company_id=company)
        self.assertEqual(policy["breaks_allowed"], 3)
        self.assertEqual(policy["break_minutes_allowed"], 20)
        self.assertEqual(policy["lunch_minutes_allowed"], 45)

    def test_activity_duration_handles_mixed_timezone_awareness(self):
        activity = SimpleNamespace(in_datetime=datetime(2026, 6, 5, 12, 0, 0))
        ended_at = timezone.make_aware(datetime(2026, 6, 5, 12, 20, 0))

        self.assertEqual(_activity_duration_seconds(activity, ended_at), 1200)

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_activity_transition_requires_verified_pin(self, _ip_allowed_mock):
        request = self.factory.post(
            "/attendance/portal/activity-transition/",
            {"employee_id": "1", "activity_type": "break", "transition": "start"},
        )
        attach_session(request)

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("PIN verification required", payload["message"])

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 0, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_start_break_closes_work_and_opens_break(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _reverse_mock,
        worked_hours_mock,
        _atomic_mock,
    ):
        request = self._verified_request()
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance = MagicMock()
        attendance_model.objects.filter.return_value.first.return_value = attendance
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "work"
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        new_activity = MagicMock()
        attendance_activity_model.objects.create.return_value = new_activity
        attendance_activity_model.objects.filter.return_value.exists.return_value = False

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["activity_type"], "break")
        self.assertEqual(payload["transition"], "start")
        self.assertEqual(open_activity.clock_out, datetime(2026, 6, 5, 12, 0, 0))
        attendance_activity_model.objects.create.assert_called_once()
        self.assertEqual(
            attendance_activity_model.objects.create.call_args.kwargs["activity_type"],
            "break",
        )
        worked_hours_mock.assert_called_once_with(employee, date(2026, 6, 5))

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 0, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_start_break_rejects_third_break(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _atomic_mock,
    ):
        request = self._verified_request()
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.activity_type = "work"
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.count.return_value = 2

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("Break limit reached", payload["message"])
        self.assertEqual(payload["breaks_taken"], 2)
        self.assertTrue(payload["break_limit_reached"])
        attendance_activity_model.objects.create.assert_not_called()

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 15, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_end_break_closes_break_and_opens_work(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _reverse_mock,
        _worked_hours_mock,
        _atomic_mock,
    ):
        request = self._verified_request({"transition": "end"})
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "break"
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        new_work_activity = MagicMock()
        attendance_activity_model.objects.create.return_value = new_work_activity
        attendance_activity_model.objects.filter.return_value.exists.return_value = False

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["active_activity_type"], "work")
        self.assertEqual(open_activity.clock_out, datetime(2026, 6, 5, 12, 15, 0))
        self.assertEqual(
            attendance_activity_model.objects.create.call_args.kwargs["activity_type"],
            "work",
        )

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 20, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_end_break_allows_overage_and_returns_minutes(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _reverse_mock,
        _worked_hours_mock,
        _atomic_mock,
    ):
        request = self._verified_request({"transition": "end"})
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "break"
        open_activity.in_datetime = datetime(2026, 6, 5, 12, 0, 0)
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.count.return_value = 1
        attendance_activity_model.objects.create.return_value = MagicMock()

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["break_overage_minutes"], 5)
        self.assertIn("exceeded", payload["message"])
        self.assertEqual(payload["active_activity_type"], "work")

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 12, 0, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_lunch_start_rejects_second_lunch(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _atomic_mock,
    ):
        request = self._verified_request(
            {"activity_type": "lunch", "transition": "start"}
        )
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.activity_type = "work"
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.exists.return_value = True

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("already", payload["message"])
        attendance_activity_model.objects.create.assert_not_called()

    @patch("attendance.views.portal.transaction.atomic", return_value=nullcontext())
    @patch("attendance.views.portal.calculate_worked_hours", return_value="03:00")
    @patch("attendance.views.portal._reverse_geocode", return_value="Test Address")
    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.get_real_now", return_value=datetime(2026, 6, 5, 13, 5, 0))
    @patch("attendance.views.portal.Attendance")
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_end_lunch_allows_overage_and_returns_minutes(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        attendance_model,
        _now_mock,
        _geo_mock,
        _reverse_mock,
        _worked_hours_mock,
        _atomic_mock,
    ):
        request = self._verified_request(
            {"activity_type": "lunch", "transition": "end"}
        )
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        attendance_model.objects.filter.return_value.first.return_value = MagicMock()
        open_activity = MagicMock()
        open_activity.attendance_date = date(2026, 6, 5)
        open_activity.shift_day = MagicMock()
        open_activity.activity_type = "lunch"
        open_activity.in_datetime = datetime(2026, 6, 5, 12, 0, 0)
        attendance_activity_model.objects.select_for_update.return_value.filter.return_value.order_by.return_value.last.return_value = open_activity
        attendance_activity_model.objects.filter.return_value.count.return_value = 0
        attendance_activity_model.objects.filter.return_value.exists.return_value = True
        attendance_activity_model.objects.create.return_value = MagicMock()

        response = public_activity_transition(request)
        payload = json.loads(response.content)

        self.assertTrue(payload["success"])
        self.assertEqual(payload["lunch_overage_minutes"], 5)
        self.assertIn("exceeded", payload["message"])
        self.assertEqual(payload["active_activity_type"], "work")

    @patch("attendance.views.portal._geofence_check", return_value=None)
    @patch("attendance.views.portal.AttendanceActivity")
    @patch("attendance.views.portal.Employee")
    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_out_rejects_active_break(
        self,
        _ip_allowed_mock,
        employee_model,
        attendance_activity_model,
        _geo_mock,
    ):
        request = self.factory.post(
            "/attendance/portal/clock-out/",
            {"employee_id": "1", "latitude": "14.6001", "longitude": "121.0001"},
        )
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }
        employee = self._employee()
        employee_model.objects.get.return_value = employee
        open_activity = MagicMock()
        open_activity.activity_type = "break"
        attendance_activity_model.objects.filter.return_value.order_by.return_value.last.return_value = open_activity

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("end your break", payload["message"])


class PortalPinGateTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_out_requires_verified_pin(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-out/", {"employee_id": "1"})
        attach_session(request)

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("PIN verification required", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_out_rejects_mismatched_pin_session(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-out/", {"employee_id": "1"})
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "2",
            "verified_at": datetime.now().timestamp(),
        }

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("does not match", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_out_allows_old_verified_pin_session(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-out/", {"employee_id": "1"})
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": (datetime.now() - timedelta(minutes=5)).timestamp(),
        }

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("Location is required", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_out_allows_flow_after_valid_pin_session(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-out/", {"employee_id": "1"})
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }

        response = public_clock_out(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("Location is required", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_in_requires_verified_pin(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-in/", {"employee_id": "1"})
        attach_session(request)

        response = public_clock_in(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("PIN verification required", payload["message"])

    @patch("attendance.views.portal._ip_is_allowed", return_value=True)
    def test_clock_in_allows_flow_after_valid_pin_session(
        self,
        _ip_allowed_mock,
    ):
        request = self.factory.post("/attendance/portal/clock-in/", {"employee_id": "1"})
        attach_session(request)
        request.session["portal_pin_verification"] = {
            "employee_id": "1",
            "verified_at": datetime.now().timestamp(),
        }

        response = public_clock_in(request)
        payload = json.loads(response.content)

        self.assertFalse(payload["success"])
        self.assertIn("Location is required", payload["message"])


class PortalUrlRoutingTests(SimpleTestCase):
    def test_portal_urls_resolve(self):
        self.assertEqual(resolve("/attendance/portal/").url_name, "public-portal")
        self.assertEqual(
            resolve("/attendance/portal/employee-lookup/").url_name,
            "portal-employee-lookup",
        )
        self.assertEqual(
            resolve("/attendance/portal/forgot-pin/").url_name,
            "portal-forgot-pin",
        )
        self.assertEqual(
            resolve("/attendance/portal/reset-pin/sample-token/").url_name,
            "portal-reset-pin",
        )
        self.assertEqual(
            resolve("/attendance/portal/verify-pin/").url_name,
            "portal-verify-pin",
        )
        self.assertEqual(
            resolve("/attendance/portal/clock-in/").url_name,
            "portal-clock-in",
        )
        self.assertEqual(
            resolve("/attendance/portal/clock-out/").url_name,
            "portal-clock-out",
        )
        self.assertEqual(
            resolve("/attendance/portal/activity-transition/").url_name,
            "portal-activity-transition",
        )
        self.assertEqual(
            resolve("/attendance/portal/attendance-history/").url_name,
            "portal-attendance-history",
        )
        self.assertEqual(
            resolve("/attendance/portal/server-time/").url_name,
            "portal-server-time",
        )

    def test_self_service_urls_do_not_resolve_after_hard_cutover(self):
        with self.assertRaises(Resolver404):
            resolve("/attendance/self-service/")
        with self.assertRaises(Resolver404):
            resolve("/attendance/self-service/employee-lookup/")
        with self.assertRaises(Resolver404):
            resolve("/attendance/self-service/clock-in/")
        with self.assertRaises(Resolver404):
            resolve("/attendance/self-service/clock-out/")
        with self.assertRaises(Resolver404):
            resolve("/attendance/self-service/server-time/")


class AttendanceDeleteTests(SimpleTestCase):
    @patch("attendance.models.HorillaModel.delete")
    @patch("attendance.models.AttendanceActivity")
    @patch("attendance.models.AttendanceLateComeEarlyOut")
    def test_delete_removes_related_late_early_and_activities(
        self,
        late_early_model,
        activity_model,
        horilla_delete_mock,
    ):
        attendance = Attendance(attendance_date=date(2026, 4, 10))
        employee = MagicMock()
        overtime_qs = MagicMock()
        overtime_qs.exists.return_value = False
        employee.employee_overtime.filter.return_value = overtime_qs
        attendance.employee_id_id = 1
        attendance._state.fields_cache["employee_id"] = employee

        attendance.delete()

        late_early_model.objects.filter.assert_called_once_with(attendance_id=attendance)
        activity_model.objects.filter.assert_called_once_with(
            attendance_date=attendance.attendance_date,
            employee_id=employee,
        )
        horilla_delete_mock.assert_called_once()

    @patch("attendance.models.HorillaModel.delete")
    @patch("attendance.models.AttendanceActivity")
    @patch("attendance.models.AttendanceLateComeEarlyOut")
    def test_delete_updates_overtime_when_available(
        self,
        _late_early_model,
        _activity_model,
        _horilla_delete_mock,
    ):
        attendance = Attendance(attendance_date=date(2026, 4, 10))
        employee = MagicMock()
        overtime_qs = MagicMock()
        overtime_qs.exists.return_value = True
        overtime_instance = MagicMock()
        overtime_qs.first.return_value = overtime_instance
        employee.employee_overtime.filter.return_value = overtime_qs
        attendance.employee_id_id = 1
        attendance._state.fields_cache["employee_id"] = employee
        attendance.update_ot = MagicMock()

        attendance.delete()

        attendance.update_ot.assert_called_once_with(overtime_instance)


class DeleteBlockedMessageTests(SimpleTestCase):
    def test_delete_blocked_message_includes_related_model_names(self):
        protected_one = MagicMock()
        protected_one._meta.verbose_name = "work record"
        protected_two = MagicMock()
        protected_two._meta.verbose_name = "late come early out"

        message = _delete_blocked_message([protected_one, protected_two])

        self.assertIn("Deletion blocked by related records:", message)
        self.assertIn("Work record", message)
        self.assertIn("Late come early out", message)


class AttendanceDeleteAuthorizationTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def _build_request(self, path):
        request = self.factory.post(path)
        employee = MagicMock()
        employee.is_active = True

        user = MagicMock()
        user.is_authenticated = True
        user.is_active = True
        user.employee_get = employee
        user.has_perm.return_value = False

        request.user = user
        request.session = {}
        return request

    @patch("employee.models.EmployeeWorkInformation.objects.filter")
    @patch("attendance.views.views.messages.success")
    @patch("attendance.views.views.HorillaRedirect", return_value=HttpResponse(status=302))
    @patch("attendance.views.views.Attendance.objects.get")
    def test_reporting_manager_can_access_attendance_delete(
        self,
        attendance_get_mock,
        _horilla_redirect_mock,
        _messages_success_mock,
        reporting_manager_filter_mock,
    ):
        request = self._build_request("/attendance/attendance-delete/1/")
        reporting_manager_filter_mock.return_value.exists.return_value = True
        attendance = MagicMock()
        attendance.attendance_date = date(2026, 4, 10)
        attendance.employee_id.employee_overtime.filter.return_value.last.return_value = None
        attendance_get_mock.return_value = attendance

        response = attendance_views.attendance_delete(request, 1)

        self.assertEqual(response.status_code, 302)
        attendance_get_mock.assert_called_once_with(id=1)

    @patch("horilla.decorators.handle_no_permission", return_value=HttpResponse(status=403))
    @patch("employee.models.EmployeeWorkInformation.objects.filter")
    @patch("attendance.views.views.Attendance.objects.get")
    def test_non_manager_without_permission_cannot_access_attendance_delete(
        self,
        attendance_get_mock,
        reporting_manager_filter_mock,
        _no_permission_mock,
    ):
        request = self._build_request("/attendance/attendance-delete/1/")
        reporting_manager_filter_mock.return_value.exists.return_value = False

        response = attendance_views.attendance_delete(request, 1)

        self.assertEqual(response.status_code, 403)
        attendance_get_mock.assert_not_called()

    @patch("employee.models.EmployeeWorkInformation.objects.filter")
    @patch("attendance.views.views.AttendanceOverTime.objects.filter")
    @patch("attendance.views.views.Attendance.objects.filter")
    def test_reporting_manager_can_access_bulk_delete(
        self,
        attendance_filter_mock,
        overtime_filter_mock,
        reporting_manager_filter_mock,
    ):
        request = self._build_request("/attendance/attendance-bulk-delete")
        request.POST = request.POST.copy()
        request.POST.setlist("ids", ["1", "2"])
        reporting_manager_filter_mock.return_value.exists.return_value = True

        attendances_qs = MagicMock()
        attendances_qs.values_list.return_value = []
        attendances_qs.__iter__.return_value = iter([])
        attendance_filter_mock.return_value = attendances_qs

        overtime_qs = MagicMock()
        overtime_qs.in_bulk.return_value = {}
        overtime_filter_mock.return_value = overtime_qs

        response = attendance_views.attendance_bulk_delete(request)

        self.assertEqual(response.status_code, 302)
        attendance_filter_mock.assert_called_once()

    @patch("employee.models.EmployeeWorkInformation.objects.filter")
    @patch("attendance.views.views.messages.success")
    @patch("attendance.views.views.HorillaRedirect", return_value=HttpResponse(status=302))
    @patch("attendance.views.views.Attendance.objects.get")
    def test_delete_still_executes_when_overtime_row_missing(
        self,
        attendance_get_mock,
        _redirect_mock,
        _messages_success_mock,
        reporting_manager_filter_mock,
    ):
        request = self._build_request("/attendance/attendance-delete/1/")
        reporting_manager_filter_mock.return_value.exists.return_value = True

        attendance = MagicMock()
        attendance.attendance_date = date(2026, 4, 10)
        attendance.attendance_overtime_approve = False
        attendance.employee_id.employee_overtime.filter.return_value.last.return_value = None
        attendance_get_mock.return_value = attendance

        response = attendance_views.attendance_delete(request, 1)

        self.assertEqual(response.status_code, 302)
        attendance.delete.assert_called_once()


class AttendanceLocationMigrationTests(SimpleTestCase):
    @staticmethod
    def _run_migration(apps):
        migration_module = importlib.import_module(
            "attendance.migrations.0005_attendanceactivity_separate_clock_locations"
        )
        migration_module.backfill_separate_clock_locations(apps, None)

    def test_backfill_closed_activity_sets_clock_out_location_only(self):
        activity = SimpleNamespace(
            clock_out=datetime(2026, 4, 10, 17, 0).time(),
            latitude=14.601,
            longitude=121.001,
            gps_address="Closed Activity Address",
            clock_in_latitude=None,
            clock_in_longitude=None,
            clock_in_gps_address=None,
            clock_out_latitude=None,
            clock_out_longitude=None,
            clock_out_gps_address=None,
            save=MagicMock(),
        )

        activities_qs = MagicMock()
        activities_qs.iterator.return_value = iter([activity])
        model = SimpleNamespace(objects=MagicMock())
        model.objects.filter.return_value = activities_qs
        apps = MagicMock()
        apps.get_model.return_value = model

        self._run_migration(apps)

        self.assertIsNone(activity.clock_in_latitude)
        self.assertIsNone(activity.clock_in_longitude)
        self.assertIsNone(activity.clock_in_gps_address)
        self.assertEqual(activity.clock_out_latitude, 14.601)
        self.assertEqual(activity.clock_out_longitude, 121.001)
        self.assertEqual(activity.clock_out_gps_address, "Closed Activity Address")
        activity.save.assert_called_once()

    def test_backfill_open_activity_sets_clock_in_location_only(self):
        activity = SimpleNamespace(
            clock_out=None,
            latitude=14.602,
            longitude=121.002,
            gps_address="Open Activity Address",
            clock_in_latitude=None,
            clock_in_longitude=None,
            clock_in_gps_address=None,
            clock_out_latitude=None,
            clock_out_longitude=None,
            clock_out_gps_address=None,
            save=MagicMock(),
        )

        activities_qs = MagicMock()
        activities_qs.iterator.return_value = iter([activity])
        model = SimpleNamespace(objects=MagicMock())
        model.objects.filter.return_value = activities_qs
        apps = MagicMock()
        apps.get_model.return_value = model

        self._run_migration(apps)

        self.assertEqual(activity.clock_in_latitude, 14.602)
        self.assertEqual(activity.clock_in_longitude, 121.002)
        self.assertEqual(activity.clock_in_gps_address, "Open Activity Address")
        self.assertIsNone(activity.clock_out_latitude)
        self.assertIsNone(activity.clock_out_longitude)
        self.assertIsNone(activity.clock_out_gps_address)
        activity.save.assert_called_once()


class AttendanceActivityUpdateViewTests(TestCase):
    def setUp(self):
        _thread_locals.request = None
        unique_id = uuid.uuid4().hex[:8]
        self.attendance_date = date(2026, 6, 1)
        self.shift_day, _created = EmployeeShiftDay.objects.get_or_create(day="monday")
        self.shift = EmployeeShift.objects.create(
            employee_shift="Activity Edit Shift",
            weekly_full_time="40:00",
            full_time="200:00",
        )
        EmployeeShiftSchedule.objects.create(
            day=self.shift_day,
            shift_id=self.shift,
            minimum_working_hour="08:00",
            start_time=time(8, 0),
            end_time=time(17, 0),
        )
        self.work_type = WorkType.objects.create(work_type="Office")
        self.employee = Employee.objects.create(
            employee_first_name="Activity",
            employee_last_name="Editor",
            email=f"activity-editor-{unique_id}@example.com",
            phone=f"0917{unique_id[:7]}",
            gender="male",
            is_active=True,
        )
        self.attendance = Attendance.objects.create(
            employee_id=self.employee,
            attendance_date=self.attendance_date,
            shift_id=self.shift,
            work_type_id=self.work_type,
            attendance_day=self.shift_day,
            attendance_clock_in_date=self.attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=self.attendance_date,
            attendance_clock_out=time(17, 0),
            attendance_worked_hour="09:00",
            minimum_hour="08:00",
            attendance_validated=True,
        )
        self.work_activity = AttendanceActivity.objects.create(
            employee_id=self.employee,
            attendance_date=self.attendance_date,
            shift_day=self.shift_day,
            clock_in_date=self.attendance_date,
            clock_in=time(8, 0),
            clock_out_date=self.attendance_date,
            clock_out=time(12, 0),
            activity_type="work",
        )
        self.break_activity = AttendanceActivity.objects.create(
            employee_id=self.employee,
            attendance_date=self.attendance_date,
            shift_day=self.shift_day,
            clock_in_date=self.attendance_date,
            clock_in=time(12, 0),
            clock_out_date=self.attendance_date,
            clock_out=time(12, 30),
            activity_type="break",
        )
        self.user = User.objects.create_user(
            username=f"activity-editor-{unique_id}",
            password="password",
        )
        Employee.objects.filter(id=self.employee.id).update(
            employee_user_id_id=self.user.id
        )
        permission = Permission.objects.get(
            codename="change_attendanceactivity",
            content_type__app_label="attendance",
        )
        self.user.user_permissions.add(permission)
        self.url = reverse(
            "attendance-activity-update",
            args=[self.employee.id, self.attendance_date.isoformat()],
        )

    def tearDown(self):
        _thread_locals.request = None

    def _post_data(self, work_out="13:00", work_out_date=None):
        if work_out_date is None:
            work_out_date = self.attendance_date.isoformat()
        return {
            "employee_id": str(self.employee.id),
            "attendance_date": self.attendance_date.isoformat(),
            "shift_id": str(self.shift.id),
            "work_type_id": str(self.work_type.id),
            "attendance_clock_in_date": self.attendance_date.isoformat(),
            "attendance_clock_in": "08:15",
            "attendance_clock_out_date": work_out_date,
            "attendance_clock_out": work_out,
            "attendance_worked_hour": "04:45",
            "minimum_hour": "08:00",
        }

    def test_get_requires_permission_and_renders_attendance_style_modal_form(self):
        self.client.force_login(self.user)

        response = self.client.get(self.url, HTTP_HX_REQUEST="true")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "attendanceActivityUpdateForm")
        self.assertContains(response, "attendance_clock_in")
        self.assertContains(response, "attendance_worked_hour")
        self.assertContains(response, "Office")
        self.assertContains(response, 'value="2026-06-01"')
        self.assertContains(response, 'background-color: #fff')
        self.assertNotContains(response, "batch_attendance_id")
        self.assertNotContains(response, "attendance_overtime_approve")
        self.assertNotContains(response, "attendance_validated")

    @patch("attendance.views.views.recalculate_attendance_for_shift")
    def test_post_updates_existing_activity_segments(self, recalculate_mock):
        self.client.force_login(self.user)

        response = self.client.post(
            f"{self.url}?page=2",
            data=self._post_data(),
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Attendance activity updated.")
        self.assertContains(
            response,
            'hx-get="/attendance/attendance-activity-search?page=2"',
        )
        self.work_activity.refresh_from_db()
        self.break_activity.refresh_from_db()
        self.assertEqual(self.work_activity.clock_in, time(8, 15))
        self.assertEqual(self.work_activity.clock_out, time(13, 0))
        self.assertEqual(self.break_activity.activity_type, "break")
        recalculate_mock.assert_called_once_with(self.shift)

    @patch("attendance.views.views.recalculate_attendance_for_shift")
    def test_post_allows_blank_checkout_fields_for_open_segment(self, recalculate_mock):
        self.client.force_login(self.user)

        response = self.client.post(
            self.url,
            data=self._post_data(work_out="", work_out_date=""),
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Attendance activity updated.")
        self.work_activity.refresh_from_db()
        self.assertIsNone(self.work_activity.clock_out_date)
        self.assertIsNone(self.work_activity.clock_out)
        recalculate_mock.assert_called_once_with(self.shift)

    @patch("attendance.views.views.recalculate_attendance_for_shift")
    def test_post_allows_activity_checkout_later_today(self, recalculate_mock):
        today = date.today()
        attendance = Attendance.objects.create(
            employee_id=self.employee,
            attendance_date=today,
            shift_id=self.shift,
            work_type_id=self.work_type,
            attendance_day=self.shift_day,
            attendance_clock_in_date=today,
            attendance_clock_in=time(0, 1),
            attendance_worked_hour="00:00",
            minimum_hour="08:00",
            attendance_validated=True,
        )
        AttendanceActivity.objects.create(
            employee_id=self.employee,
            attendance_date=today,
            shift_day=self.shift_day,
            clock_in_date=today,
            clock_in=time(0, 1),
            activity_type="work",
        )
        url = reverse(
            "attendance-activity-update",
            args=[self.employee.id, today.isoformat()],
        )
        post_data = self._post_data(work_out="23:59", work_out_date=today.isoformat())
        post_data.update(
            {
                "attendance_date": today.isoformat(),
                "attendance_clock_in_date": today.isoformat(),
                "attendance_clock_in": "00:01",
                "attendance_worked_hour": "23:58",
            }
        )
        self.client.force_login(self.user)

        response = self.client.post(url, data=post_data, HTTP_HX_REQUEST="true")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Check-out time cannot be in the future")
        self.assertContains(response, "Attendance activity updated.")
        attendance.refresh_from_db()
        self.assertEqual(attendance.attendance_clock_out, time(23, 59))
        recalculate_mock.assert_called_once_with(self.shift)

    @patch("attendance.views.views.recalculate_attendance_for_shift")
    def test_invalid_checkout_before_checkin_does_not_save(self, recalculate_mock):
        self.client.force_login(self.user)

        response = self.client.post(
            self.url,
            data=self._post_data(work_out="07:30"),
            HTTP_HX_REQUEST="true",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Check out cannot be earlier than check in.")
        self.work_activity.refresh_from_db()
        self.assertEqual(self.work_activity.clock_in, time(8, 0))
        self.assertEqual(self.work_activity.clock_out, time(12, 0))
        recalculate_mock.assert_not_called()

    def test_user_without_permission_is_blocked(self):
        unique_id = uuid.uuid4().hex[:8]
        user = User.objects.create_user(
            username=f"activity-viewer-{unique_id}",
            password="password",
        )
        viewer = Employee.objects.create(
            employee_first_name="Activity",
            employee_last_name="Viewer",
            email=f"activity-viewer-{unique_id}@example.com",
            phone=f"0918{unique_id[:7]}",
            gender="male",
            is_active=True,
        )
        Employee.objects.filter(id=viewer.id).update(employee_user_id_id=user.id)
        self.client.force_login(user)

        response = self.client.get(self.url, HTTP_HX_REQUEST="true")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "attendanceActivityUpdateForm")


class AttendanceActivityMetaBuilderTests(SimpleTestCase):
    @patch("attendance.views.views.AttendanceActivity")
    def test_build_meta_returns_separate_in_and_out_location_data(
        self,
        attendance_activity_model,
    ):
        attendance_row = SimpleNamespace(
            id=1,
            employee_id_id=101,
            attendance_date=date(2026, 4, 10),
        )
        paginated_attendances = SimpleNamespace(object_list=[attendance_row])

        activities = [
            SimpleNamespace(
                employee_id_id=101,
                attendance_date=date(2026, 4, 10),
                clock_in_selfie=None,
                clock_out_selfie=None,
                clock_in_gps_address="Clock In Address",
                clock_out_gps_address=None,
                clock_in_latitude=14.501,
                clock_in_longitude=120.901,
                clock_out_latitude=None,
                clock_out_longitude=None,
                gps_address="Legacy In Address",
                latitude=14.501,
                longitude=120.901,
                clock_out=None,
            ),
            SimpleNamespace(
                employee_id_id=101,
                attendance_date=date(2026, 4, 10),
                clock_in_selfie=None,
                clock_out_selfie=None,
                clock_in_gps_address=None,
                clock_out_gps_address="Clock Out Address",
                clock_in_latitude=None,
                clock_in_longitude=None,
                clock_out_latitude=14.601,
                clock_out_longitude=121.001,
                gps_address="Legacy Out Address",
                latitude=14.601,
                longitude=121.001,
                clock_out=datetime(2026, 4, 10, 17, 0).time(),
            ),
        ]

        attendance_activity_model.objects.filter.return_value.order_by.return_value = (
            activities
        )

        result = build_my_attendance_activity_meta(paginated_attendances)
        meta = result[1]

        self.assertEqual(meta["check_in_location"], "Clock In Address")
        self.assertEqual(meta["check_out_location"], "Clock Out Address")
        self.assertEqual(
            meta["check_in_maps_url"],
            "https://www.google.com/maps?q=14.501,120.901",
        )
        self.assertEqual(
            meta["check_out_maps_url"],
            "https://www.google.com/maps?q=14.601,121.001",
        )
        self.assertEqual(meta["location"], "Clock Out Address")
        self.assertEqual(
            meta["maps_url"],
            "https://www.google.com/maps?q=14.601,121.001",
        )


class FakeAttendanceQuerySet:
    def __init__(self):
        self.filters = []
        self.selected_related = []

    def filter(self, **kwargs):
        self.filters.append(kwargs)
        return self

    def select_related(self, *fields):
        self.selected_related.extend(fields)
        return self


class AttendanceLazyTabContextTests(SimpleTestCase):
    def _request(self, querystring):
        request = RequestFactory().get(f"/attendance/attendance-search?{querystring}")
        request.user = SimpleNamespace(has_perm=lambda perm: True)
        return request

    @patch("attendance.views.views.get_key_instances")
    @patch("attendance.views.views.build_daily_attendance_rows")
    @patch("attendance.views.views.paginator_qry")
    @patch("attendance.views.views.filtersubordinates")
    @patch("attendance.views.views.AttendanceFilters")
    @patch("attendance.views.views.AttendanceValidationCondition")
    @patch("attendance.views.views.Attendance")
    def test_overtime_tab_uses_overtime_queryset_and_page_param(
        self,
        attendance_model,
        validation_condition,
        attendance_filters,
        filter_subordinates,
        paginator,
        build_rows,
        _get_key_instances,
    ):
        queryset = FakeAttendanceQuerySet()
        attendance_model.objects.filter.return_value = queryset
        validation_condition.objects.first.return_value = None
        attendance_filters.side_effect = (
            lambda data, queryset=None, **kwargs: SimpleNamespace(qs=queryset)
        )
        filter_subordinates.side_effect = lambda request, qs, perm: qs
        page = SimpleNamespace(object_list=[SimpleNamespace(id=7)])
        paginator.return_value = page
        build_rows.return_value = [SimpleNamespace(row_key="attendance-7")]

        context = build_attendance_tab_context(
            self._request("tab=overtime&opage=2")
        )

        self.assertIn(
            {"overtime_second__gt": 0, "attendance_validated": True},
            queryset.filters,
        )
        paginator.assert_called_once_with(queryset, "2")
        self.assertEqual(context["active_tab_key"], "overtime")
        self.assertEqual(context["overtime_attendances"], page)
        self.assertEqual(context["ot_attendances_ids"], "[7]")

    @patch("attendance.views.views.get_key_instances")
    @patch("attendance.views.views.build_daily_attendance_rows")
    @patch("attendance.views.views.group_by_queryset")
    @patch("attendance.views.views.filtersubordinates")
    @patch("attendance.views.views.AttendanceFilters")
    @patch("attendance.views.views.AttendanceValidationCondition")
    @patch("attendance.views.views.Attendance")
    def test_validated_group_by_uses_validated_page_param_only(
        self,
        attendance_model,
        validation_condition,
        attendance_filters,
        filter_subordinates,
        group_by,
        build_rows,
        _get_key_instances,
    ):
        queryset = FakeAttendanceQuerySet()
        attendance_model.objects.filter.return_value = queryset
        validation_condition.objects.first.return_value = None
        attendance_filters.side_effect = (
            lambda data, queryset=None, **kwargs: SimpleNamespace(qs=queryset)
        )
        filter_subordinates.side_effect = lambda request, qs, perm: qs
        grouped_page = SimpleNamespace(
            object_list=[
                {
                    "list": SimpleNamespace(
                        object_list=[SimpleNamespace(id=3)],
                    )
                }
            ]
        )
        group_by.return_value = grouped_page
        build_rows.return_value = [SimpleNamespace(row_key="attendance-3")]

        context = build_attendance_tab_context(
            self._request("tab=validated&page=4&field=employee_id")
        )

        self.assertIn({"attendance_validated": True}, queryset.filters)
        group_by.assert_called_once_with(queryset, "employee_id", "4", "page")
        self.assertEqual(context["active_tab_key"], "validated")
        self.assertTrue(context["is_grouped"])
        self.assertEqual(context["attendances_ids"], "[3]")


class FakeActivityQuerySet(list):
    def select_related(self, *args):
        return self


class AttendanceExportJobEndpointTests(TestCase):
    def setUp(self):
        _thread_locals.request = None
        self.media_dir = tempfile.mkdtemp()
        self.settings_override = override_settings(MEDIA_ROOT=self.media_dir)
        self.settings_override.enable()
        self.user = self._make_user("exporter", "exporter@example.com")
        permission = Permission.objects.get(
            content_type__app_label="attendance",
            codename="change_attendanceactivity",
        )
        self.user.user_permissions.add(permission)
        self.other_user = self._make_user("other-exporter", "other@example.com")

    def tearDown(self):
        _thread_locals.request = None
        self.settings_override.disable()
        shutil.rmtree(self.media_dir, ignore_errors=True)

    def _make_user(self, username, email):
        user = User.objects.create_user(
            username=username,
            password="password",
            email=email,
        )
        Employee.objects.create(
            employee_user_id=user,
            employee_first_name=username,
            email=email,
            phone="1234567890",
        )
        return user

    def test_starting_activity_export_job_returns_job_id(self):
        self.client.force_login(self.user)
        with patch("attendance.views.views.start_export_job") as start_job:
            response = self.client.post(
                reverse("attendance-activity-export-start"),
                data={"selected_fields": ["employee_id", "attendance_date"]},
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["job_id"])
        self.assertTrue(file_path(payload["job_id"]).parent.exists())
        start_job.assert_called_once()

    def test_progress_endpoint_only_returns_jobs_for_owner(self):
        job_id = create_export_job(self.user.id, "export.xlsx")
        update_job(job_id, status="running", percent=40, message="Working")

        self.client.force_login(self.other_user)
        response = self.client.get(reverse("attendance-export-progress", args=[job_id]))

        self.assertEqual(response.status_code, 404)

    def test_download_endpoint_serves_completed_owner_job(self):
        job_id = create_export_job(self.user.id, "export.xlsx")
        file_path(job_id).write_bytes(b"fake-xlsx")
        update_job(job_id, status="complete", percent=100, message="Ready")

        self.client.force_login(self.user)
        response = self.client.get(reverse("attendance-export-download", args=[job_id]))

        self.assertEqual(response.status_code, 200)
        self.assertIn("export.xlsx", response["Content-Disposition"])

    def test_download_endpoint_rejects_non_owner_and_incomplete_jobs(self):
        job_id = create_export_job(self.user.id, "export.xlsx")
        file_path(job_id).write_bytes(b"fake-xlsx")
        update_job(job_id, status="failed", percent=100, message="Failed")

        self.client.force_login(self.other_user)
        owner_response = self.client.get(
            reverse("attendance-export-download", args=[job_id])
        )
        self.assertEqual(owner_response.status_code, 404)

        self.client.force_login(self.user)
        incomplete_response = self.client.get(
            reverse("attendance-export-download", args=[job_id])
        )
        self.assertEqual(incomplete_response.status_code, 409)

    def test_failed_job_progress_exposes_error_status(self):
        job_id = create_export_job(self.user.id, "export.xlsx")
        update_job(
            job_id,
            status="failed",
            percent=100,
            message="Export failed",
            error="boom",
        )

        self.client.force_login(self.user)
        response = self.client.get(reverse("attendance-export-progress", args=[job_id]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "failed")
        self.assertEqual(payload["error"], "boom")


class AttendanceActivityExportTests(SimpleTestCase):
    def _format_export_value(self, value, employee):
        if isinstance(value, time):
            return value.strftime("%H:%M")
        return value

    def _segment(
        self,
        clock_in=None,
        clock_out=None,
        clock_in_selfie=None,
        clock_out_selfie=None,
        clock_in_location="",
        clock_out_location="",
        clock_in_date=None,
        clock_out_date=None,
        in_datetime=None,
        out_datetime=None,
    ):
        return SimpleNamespace(
            clock_in=clock_in,
            clock_out=clock_out,
            clock_in_selfie=clock_in_selfie,
            clock_out_selfie=clock_out_selfie,
            clock_in_location=clock_in_location,
            clock_out_location=clock_out_location,
            clock_in_date=clock_in_date,
            clock_out_date=clock_out_date,
            in_datetime=in_datetime,
            out_datetime=out_datetime,
        )

    def _daily_row(self, late_come_duration="00:05", early_out_duration="00:10"):
        employee = SimpleNamespace(
            id=101,
            employee_work_info=SimpleNamespace(
                branch_id="Main Branch",
                department_id="HR",
            ),
        )
        work_segment = self._segment(
            clock_in=time(9, 0),
            clock_out=time(17, 0),
            clock_in_date=date(2026, 4, 10),
            clock_out_date=date(2026, 4, 10),
        )
        return SimpleNamespace(
            employee=employee,
            attendance_date=date(2026, 4, 10),
            work_segments=[work_segment],
            work_in=self._segment(
                clock_in=time(9, 0),
                clock_in_selfie=SimpleNamespace(url="/media/clock-in.jpg"),
                clock_in_location="Clock In Address",
            ),
            work_out=self._segment(
                clock_out=time(17, 0),
                clock_out_selfie=SimpleNamespace(url="/media/clock-out.jpg"),
                clock_out_location="Clock Out Address",
            ),
            break_segments=[
                self._segment(
                    time(10, 0),
                    time(10, 15),
                    clock_in_location="Break In Address 1",
                    clock_out_location="Break Out Address 1",
                ),
                self._segment(
                    time(15, 0),
                    time(15, 10),
                    clock_in_location="Break In Address 2",
                    clock_out_location="Break Out Address 2",
                ),
            ],
            lunch_segments=[
                self._segment(
                    time(12, 0),
                    time(13, 0),
                    SimpleNamespace(url="/media/lunch-in.jpg"),
                    SimpleNamespace(url="/media/lunch-out.jpg"),
                    "Lunch In Address",
                    "Lunch Out Address",
                )
            ],
            shift="Morning",
            shift_day="friday",
            shift_schedule="Mon, Tue, Wed, Thu, Fri",
            has_shift_schedule=True,
            schedule=SimpleNamespace(
                start_time=time(9, 0), end_time=time(18, 0), is_night_shift=False
            ),
            late_come_duration=late_come_duration,
            early_out_duration=early_out_duration,
            work_hours="07:30",
            break_hours="00:25",
            lunch_hours="01:00",
            overtime="00:30",
            holiday=None,
            holiday_type=None,
            is_rest_day=False,
            is_leave_only=False,
        )

    def test_activity_export_defaults_use_split_daily_columns(self):
        form = AttendanceActivityExportForm()

        expected_fields = AttendanceActivityExportForm.default_fields
        choice_labels = {
            field_name: str(label)
            for field_name, label in form.fields["selected_fields"].choices
        }
        expected_headers = [choice_labels[field_name] for field_name in expected_fields]

        self.assertEqual(form.fields["selected_fields"].initial, expected_fields)
        self.assertEqual(
            [str(label) for _, label in _attendance_activity_export_columns(form, expected_fields)],
            expected_headers,
        )
        self.assertEqual(
            expected_fields.index("daily_basic_hours"),
            expected_fields.index("daily_work_hours") + 1,
        )
        self.assertEqual(
            expected_fields.index("daily_holiday_type"),
            expected_fields.index("daily_holiday") + 1,
        )
        self.assertEqual(
            expected_fields.index("daily_rest_day"),
            expected_fields.index("daily_shift_end") + 1,
        )
        self.assertEqual(
            expected_fields.index("daily_worked_special_holiday"),
            expected_fields.index("daily_holiday_type") + 1,
        )
        self.assertEqual(choice_labels["daily_holiday_type"], "Holiday Type")
        self.assertEqual(choice_labels["daily_rest_day"], "Rest Day")
        self.assertNotIn("daily_shift_schedule", expected_fields)
        premium_labels = {
            "daily_worked_special_holiday": "Worked on Special Holiday",
            "daily_ot_worked_special_holiday": "OT on Worked on Special Holiday",
            "daily_worked_regular_holiday": "Worked on Regular Holiday",
            "daily_ot_worked_regular_holiday": "OT on Worked on Regular Holiday",
            "daily_worked_rest_day": "Worked on Restday",
            "daily_ot_worked_rest_day": "OT on Worked Restday",
            "daily_night_differential": "Night Differential Hours",
            "daily_night_differential_overtime": "Night Differential Hours- OVERTIME",
            "daily_night_differential_rest_day_overtime": "Night Differential - Rest Day Overtime",
            "daily_night_differential_rest_day": "Night Differential Hours-REST DAY",
            "daily_night_differential_regular_holiday": "Night Differential Regular Holiday - Hours",
            "daily_night_differential_special_holiday": "Night Differential Special Holiday - Hours",
            "daily_night_differential_special_holiday_overtime": "Night Differential Hours-SPECIAL HOL. OVERTIME",
            "daily_night_differential_regular_holiday_overtime": "Night Differential - Overtime - Legal Hours",
            "daily_rest_day_regular_holiday": "Rest Day Hours- Regular Holiday",
            "daily_ot_regular_holiday_rest_day": "Overtime hours - Regular Holiday - Rest Day",
            "daily_night_differential_rest_day_regular_holiday_overtime": "Night Differential Hours-REST DAY Regular HOL. OVERTIME",
            "daily_night_differential_rest_day_regular_holiday": "Night Differential Hours-REST DAY Regular Pay.",
            "daily_rest_day_special_holiday": "Rest Day Hours - Special Holiday",
            "daily_night_differential_rest_day_special_holiday": "Night Differential Hours-REST DAY SPECIAL HOL.",
            "daily_night_differential_rest_day_special_holiday_overtime": "Night Differential Hours-REST DAY Special HOL. OVERTIME",
            "daily_ot_special_holiday_rest_day": "Overtime hours - Special Holiday - Rest Day",
        }
        for field_name, label in premium_labels.items():
            self.assertIn(field_name, expected_fields)
            self.assertEqual(choice_labels[field_name], label)

    def test_payroll_group_export_includes_premium_columns(self):
        selected_field_names = [
            field_name for field_name, _ in _payroll_group_export_selected_columns()
        ]

        self.assertIn("daily_worked_special_holiday", selected_field_names)
        self.assertIn("daily_holiday_type", selected_field_names)
        self.assertIn("daily_rest_day", selected_field_names)
        self.assertEqual(
            selected_field_names.index("daily_rest_day"),
            selected_field_names.index("daily_shift_end") + 1,
        )
        self.assertNotIn("daily_shift_schedule", selected_field_names)
        self.assertIn("daily_ot_special_holiday_rest_day", selected_field_names)
        self.assertIn(
            "daily_night_differential_rest_day_regular_holiday_overtime",
            selected_field_names,
        )

    @patch("attendance.views.views.format_export_value")
    def test_daily_export_values_split_times_and_images(self, format_export_value):
        format_export_value.side_effect = self._format_export_value
        row = self._daily_row()

        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_clock_in", None),
            "09:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_clock_out", None),
            "17:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_break_in", None),
            "10:00; 15:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_break_out", None),
            "10:15; 15:10",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_lunch_in", None),
            "12:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_lunch_out", None),
            "13:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_clock_in_location", None
            ),
            "Clock In Address",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_clock_out_location", None
            ),
            "Clock Out Address",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_break_in_location", None
            ),
            "Break In Address 1; Break In Address 2",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_break_out_location", None
            ),
            "Break Out Address 1; Break Out Address 2",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_lunch_in_location", None
            ),
            "Lunch In Address",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                row, "daily_lunch_out_location", None
            ),
            "Lunch Out Address",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_clock_image", None),
            "In /media/clock-in.jpg / Out /media/clock-out.jpg",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_lunch_image", None),
            "In /media/lunch-in.jpg / Out /media/lunch-out.jpg",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_late_come", None),
            "00:05",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_early_out", None),
            "00:10",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                self._daily_row(early_out_duration=""), "daily_early_out", None
            ),
            "00:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                self._daily_row(early_out_duration=None), "daily_early_out", None
            ),
            "00:00",
        )
        leave_row = self._daily_row(early_out_duration="")
        leave_row.is_leave_only = True
        self.assertEqual(
            _attendance_activity_daily_export_value(
                leave_row, "daily_early_out", None
            ),
            "",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_overtime", None),
            "00:30",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_basic_hours", None),
            "",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_rest_day", None),
            "",
        )
        row.holiday = "Founding Day"
        row.holiday_type = "special"
        row.is_rest_day = True
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_shift_start", None),
            "09:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_shift_end", None),
            "18:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_rest_day", None),
            "Yes",
        )
        leave_row.is_rest_day = True
        self.assertEqual(
            _attendance_activity_daily_export_value(
                leave_row, "daily_rest_day", None
            ),
            "",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_holiday_type", None),
            "Special Holiday",
        )
        row.holiday_type = "unclassified"
        row.is_rest_day = False
        self.assertEqual(
            _attendance_activity_daily_export_value(row, "daily_holiday_type", None),
            "Unclassified",
        )

    def _premium_row(
        self,
        start,
        end,
        schedule_start=time(9, 0),
        schedule_end=time(17, 0),
        holiday_type=None,
        is_rest_day=False,
        is_night_shift=False,
    ):
        row = self._daily_row()
        row.attendance_date = start.date()
        row.schedule = (
            SimpleNamespace(
                start_time=schedule_start,
                end_time=schedule_end,
                is_night_shift=is_night_shift,
            )
            if schedule_end
            else None
        )
        row.work_segments = [
            self._segment(
                clock_in=start.time(),
                clock_out=end.time(),
                clock_in_date=start.date(),
                clock_out_date=end.date(),
                in_datetime=start,
                out_datetime=end,
            )
        ]
        row.holiday_type = holiday_type
        row.is_rest_day = is_rest_day
        return row

    def _premium_value(self, row, field_name):
        return _attendance_activity_daily_export_value(row, field_name, None)

    def test_premium_export_buckets_split_holiday_rest_day_and_overtime(self):
        start = datetime(2026, 4, 10, 9, 0)
        end = datetime(2026, 4, 10, 19, 0)

        special = self._premium_row(start, end, holiday_type="special")
        self.assertEqual(self._premium_value(special, "daily_worked_special_holiday"), "08:00")
        self.assertEqual(self._premium_value(special, "daily_ot_worked_special_holiday"), "02:00")

        regular = self._premium_row(start, end, holiday_type="regular")
        self.assertEqual(self._premium_value(regular, "daily_worked_regular_holiday"), "08:00")
        self.assertEqual(self._premium_value(regular, "daily_ot_worked_regular_holiday"), "02:00")

        rest_day = self._premium_row(start, end, is_rest_day=True)
        self.assertEqual(self._premium_value(rest_day, "daily_worked_rest_day"), "08:00")
        self.assertEqual(self._premium_value(rest_day, "daily_ot_worked_rest_day"), "02:00")

        regular_rest = self._premium_row(
            start, end, holiday_type="regular", is_rest_day=True
        )
        self.assertEqual(
            self._premium_value(regular_rest, "daily_rest_day_regular_holiday"),
            "08:00",
        )
        self.assertEqual(
            self._premium_value(regular_rest, "daily_ot_regular_holiday_rest_day"),
            "02:00",
        )
        self.assertEqual(
            self._premium_value(regular_rest, "daily_worked_regular_holiday"),
            "00:00",
        )
        self.assertEqual(
            self._premium_value(regular_rest, "daily_worked_rest_day"),
            "00:00",
        )

        special_rest = self._premium_row(
            start, end, holiday_type="special", is_rest_day=True
        )
        self.assertEqual(
            self._premium_value(special_rest, "daily_rest_day_special_holiday"),
            "08:00",
        )
        self.assertEqual(
            self._premium_value(special_rest, "daily_ot_special_holiday_rest_day"),
            "02:00",
        )

    def test_premium_export_buckets_split_night_differential(self):
        row = self._premium_row(
            datetime(2026, 4, 10, 20, 0),
            datetime(2026, 4, 11, 6, 0),
            schedule_start=time(20, 0),
            schedule_end=time(4, 0),
            is_night_shift=True,
        )

        self.assertEqual(self._premium_value(row, "daily_night_differential"), "06:00")
        self.assertEqual(
            self._premium_value(row, "daily_night_differential_overtime"),
            "02:00",
        )

    def test_premium_export_buckets_split_holiday_rest_day_night_differential(self):
        row = self._premium_row(
            datetime(2026, 4, 10, 20, 0),
            datetime(2026, 4, 11, 6, 0),
            schedule_start=time(20, 0),
            schedule_end=time(4, 0),
            holiday_type="regular",
            is_rest_day=True,
            is_night_shift=True,
        )

        self.assertEqual(
            self._premium_value(
                row, "daily_night_differential_rest_day_regular_holiday"
            ),
            "06:00",
        )
        self.assertEqual(
            self._premium_value(
                row,
                "daily_night_differential_rest_day_regular_holiday_overtime",
            ),
            "02:00",
        )

        special_row = self._premium_row(
            datetime(2026, 4, 10, 21, 0),
            datetime(2026, 4, 11, 2, 0),
            schedule_start=time(21, 0),
            schedule_end=time(1, 0),
            holiday_type="special",
            is_rest_day=True,
            is_night_shift=True,
        )
        self.assertEqual(
            self._premium_value(
                special_row, "daily_night_differential_rest_day_special_holiday"
            ),
            "03:00",
        )
        self.assertEqual(
            self._premium_value(
                special_row,
                "daily_night_differential_rest_day_special_holiday_overtime",
            ),
            "01:00",
        )

    def test_premium_export_no_schedule_treats_all_work_as_non_overtime(self):
        row = self._premium_row(
            datetime(2026, 4, 10, 9, 0),
            datetime(2026, 4, 10, 19, 0),
            schedule_end=None,
            is_rest_day=True,
        )

        self.assertEqual(self._premium_value(row, "daily_worked_rest_day"), "10:00")
        self.assertEqual(self._premium_value(row, "daily_ot_worked_rest_day"), "00:00")

    def test_premium_export_blank_schedule_treats_all_work_as_non_overtime_rest_day(self):
        row = self._premium_row(
            datetime(2026, 4, 10, 9, 0),
            datetime(2026, 4, 10, 19, 0),
            is_rest_day=True,
        )
        row.schedule = SimpleNamespace(
            start_time=None,
            end_time=None,
            is_night_shift=False,
        )

        self.assertEqual(self._premium_value(row, "daily_worked_rest_day"), "10:00")
        self.assertEqual(self._premium_value(row, "daily_ot_worked_rest_day"), "00:00")

    @patch("attendance.views.views.format_export_value")
    @patch("attendance.views.views._attendance_activity_export_daily_rows")
    def test_export_data_collapses_activity_rows_to_daily_rows(
        self, daily_rows, format_export_value
    ):
        format_export_value.side_effect = self._format_export_value
        employee = SimpleNamespace(id=101)
        attendance_date = date(2026, 4, 10)
        daily_rows.return_value = [
            self._daily_row(late_come_duration="07:00", early_out_duration="02:00")
        ]
        activities = FakeActivityQuerySet(
            [
                SimpleNamespace(employee_id_id=101, attendance_date=attendance_date),
                SimpleNamespace(employee_id_id=101, attendance_date=attendance_date),
            ]
        )

        data = _attendance_activity_export_data(
            activities,
            [
                ("daily_clock_in", "Clock In"),
                ("daily_clock_out", "Clock Out"),
                ("daily_break_in", "Break In"),
                ("daily_break_out", "Break Out"),
                ("daily_break_in_location", "Break In Location"),
                ("daily_late_come", "Late Come"),
                ("daily_early_out", "Early Out"),
            ],
            employee,
        )

        self.assertEqual(data["Clock In"], ["09:00"])
        self.assertEqual(data["Clock Out"], ["17:00"])
        self.assertEqual(data["Break In"], ["10:00; 15:00"])
        self.assertEqual(data["Break Out"], ["10:15; 15:10"])
        self.assertEqual(
            data["Break In Location"],
            ["Break In Address 1; Break In Address 2"],
        )
        self.assertEqual(data["Late Come"], ["07:00"])
        self.assertEqual(data["Early Out"], ["02:00"])

    @patch("attendance.views.views.format_export_value")
    @patch("attendance.views.views._attendance_activity_export_daily_rows")
    def test_payroll_group_export_fills_every_employee_date_with_empty_rows(
        self, daily_rows, format_export_value
    ):
        format_export_value.side_effect = self._format_export_value

        def employee(pk, employee_no):
            return SimpleNamespace(
                id=pk,
                employee_no=employee_no,
                get_full_name=lambda: f"Employee {employee_no}",
                employee_work_info=SimpleNamespace(
                    branch_id="Main Branch",
                    department_id="HR",
                    payroll_group_id="Semi Monthly",
                    shift_id="Morning",
                    work_type_id="Office",
                ),
            )

        emp1 = employee(101, "E001")
        emp2 = employee(102, "E002")
        real_row = self._daily_row()
        real_row.employee = emp1
        real_row.attendance_date = date(2026, 7, 1)
        real_row.work_hours = "08:00"
        leave_row = self._daily_row(early_out_duration="")
        leave_row.employee = emp1
        leave_row.attendance_date = date(2026, 7, 2)
        leave_row.work_in = None
        leave_row.work_out = None
        leave_row.work_hours = None
        leave_row.break_hours = None
        leave_row.lunch_hours = None
        leave_row.overtime = None
        leave_row.leave = "Vacation Leave"
        leave_row.leave_type = "Vacation Leave"
        leave_row.leave_days = 1
        leave_row.is_leave_only = True
        daily_rows.return_value = [real_row, leave_row]

        rows = _attendance_activity_payroll_group_daily_rows(
            FakeActivityQuerySet([]),
            [emp1, emp2],
            date(2026, 7, 1),
            date(2026, 7, 3),
        )
        data = _attendance_activity_export_data_from_rows(
            rows,
            [
                ("employee_number", "Employee No."),
                ("attendance_date", "Attendance Date"),
                ("daily_clock_in", "Clock In"),
                ("daily_early_out", "Early Out"),
                ("daily_work_hours", "Work Hours"),
                ("daily_basic_hours", "Basic Hours"),
                ("daily_break_hours", "Break Hours"),
                ("daily_lunch_hours", "Lunch Hours"),
                ("daily_overtime", "Overtime"),
                ("daily_worked_special_holiday", "Worked on Special Holiday"),
            ],
            None,
        )

        self.assertEqual(len(rows), 6)
        self.assertEqual(
            data["Employee No."],
            ["E001", "E001", "E001", "E002", "E002", "E002"],
        )
        self.assertEqual(
            data["Attendance Date"],
            [
                date(2026, 7, 1),
                date(2026, 7, 2),
                date(2026, 7, 3),
                date(2026, 7, 1),
                date(2026, 7, 2),
                date(2026, 7, 3),
            ],
        )
        self.assertEqual(data["Clock In"], ["09:00", "", "", "", "", ""])
        self.assertEqual(
            data["Early Out"], ["00:10", "", "00:00", "00:00", "00:00", "00:00"]
        )
        self.assertEqual(
            data["Work Hours"], ["08:00", "", "00:00", "00:00", "00:00", "00:00"]
        )
        self.assertEqual(data["Basic Hours"], ["", "", "", "", "", ""])
        self.assertEqual(
            data["Break Hours"], ["00:25", "", "00:00", "00:00", "00:00", "00:00"]
        )
        self.assertEqual(
            data["Lunch Hours"], ["01:00", "", "00:00", "00:00", "00:00", "00:00"]
        )
        self.assertEqual(
            data["Overtime"], ["00:30", "", "00:00", "00:00", "00:00", "00:00"]
        )
        self.assertEqual(
            data["Worked on Special Holiday"],
            ["00:00", "", "00:00", "00:00", "00:00", "00:00"],
        )

    @patch("attendance.views.views.format_export_value")
    @patch("attendance.views.views._attendance_activity_export_daily_rows")
    def test_payroll_group_export_keeps_rest_day_ot_from_fallback_schedule(
        self, daily_rows, format_export_value
    ):
        format_export_value.side_effect = self._format_export_value
        attendance_date = date(2026, 7, 4)
        employee = SimpleNamespace(
            id=101,
            employee_no="E001",
            get_full_name=lambda: "Employee E001",
            employee_work_info=SimpleNamespace(
                branch_id="Main Branch",
                department_id="HR",
                payroll_group_id="Semi Monthly",
                shift_id="Morning",
                work_type_id="Office",
            ),
        )
        row = self._premium_row(
            datetime(2026, 7, 4, 9, 0),
            datetime(2026, 7, 4, 18, 0),
            schedule_start=time(8, 0),
            schedule_end=time(17, 0),
            is_rest_day=True,
        )
        row.employee = employee
        row.attendance_date = attendance_date
        row.work_in = self._segment(
            clock_in=time(9, 0),
            clock_in_date=attendance_date,
        )
        row.work_out = self._segment(
            clock_out=time(18, 0),
            clock_out_date=attendance_date,
        )
        daily_rows.return_value = [row]

        rows = _attendance_activity_payroll_group_daily_rows(
            FakeActivityQuerySet(
                [
                    SimpleNamespace(
                        employee_id_id=employee.id,
                        attendance_date=attendance_date,
                    )
                ]
            ),
            [employee],
            attendance_date,
            attendance_date,
        )
        data = _attendance_activity_export_data_from_rows(
            rows,
            [
                ("employee_number", "Employee No."),
                ("attendance_date", "Attendance Date"),
                ("daily_shift_start", "Shift Start"),
                ("daily_shift_end", "Shift End"),
                ("daily_rest_day", "Rest Day"),
                ("daily_worked_rest_day", "Worked on Restday"),
                ("daily_ot_worked_rest_day", "OT on Worked Restday"),
            ],
            None,
        )

        self.assertEqual(data["Shift Start"], ["08:00"])
        self.assertEqual(data["Shift End"], ["17:00"])
        self.assertEqual(data["Rest Day"], ["Yes"])
        self.assertNotIn("__Rest Day", data)
        self.assertEqual(data["Worked on Restday"], ["08:00"])
        self.assertEqual(data["OT on Worked Restday"], ["01:00"])

    def _formula_workbook(
        self, df, selected_columns, total_ranges=None, totals_sheet_name=None
    ):
        output = io.BytesIO()
        writer = pd.ExcelWriter(output, engine="xlsxwriter")
        writer.book.set_calc_mode("auto")
        df.to_excel(writer, index=False, sheet_name="Sheet1")
        _write_export_row_formulas(
            writer.book,
            writer.sheets["Sheet1"],
            df,
            selected_columns,
            total_ranges or [],
        )
        if totals_sheet_name:
            _write_export_totals_sheet(
                writer,
                "Sheet1",
                df,
                selected_columns,
                total_ranges or [],
                totals_sheet_name,
            )
        writer.close()
        output.seek(0)
        return load_workbook(output, data_only=False)

    def _cell_formula_text(self, cell):
        return getattr(cell.value, "text", cell.value)

    def test_export_row_duration_columns_are_excel_formulas(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_break_in", "Break In"),
            ("daily_break_out", "Break Out"),
            ("daily_lunch_in", "Lunch In"),
            ("daily_lunch_out", "Lunch Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_late_come", "Late Come"),
            ("daily_early_out", "Early Out"),
            ("daily_work_hours", "Work Hours"),
            ("daily_basic_hours", "Basic Hours"),
            ("daily_break_hours", "Break Hours"),
            ("daily_lunch_hours", "Lunch Hours"),
            ("daily_overtime", "Overtime"),
            ("daily_leave_type", "Leave Type"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E001"],
                "Clock In": ["09:15", ""],
                "Clock Out": ["18:30", ""],
                "Break In": ["10:00; 15:00", ""],
                "Break Out": ["10:15; 15:10", ""],
                "Lunch In": ["12:00", ""],
                "Lunch Out": ["13:00", ""],
                "Shift Start": ["09:00", ""],
                "Shift End": ["18:00", ""],
                "Late Come": ["", ""],
                "Early Out": ["", ""],
                "Work Hours": ["", ""],
                "Basic Hours": ["", ""],
                "Break Hours": ["", ""],
                "Lunch Hours": ["", ""],
                "Overtime": ["", ""],
                "Leave Type": ["", "Vacation Leave"],
            }
        )

        workbook = self._formula_workbook(df, selected_columns)
        sheet = workbook["Sheet1"]

        late_formula = self._cell_formula_text(sheet["J2"])
        early_formula = self._cell_formula_text(sheet["K2"])
        work_formula = self._cell_formula_text(sheet["L2"])
        basic_formula = self._cell_formula_text(sheet["M2"])
        break_formula = self._cell_formula_text(sheet["N2"])
        lunch_formula = self._cell_formula_text(sheet["O2"])
        overtime_formula = self._cell_formula_text(sheet["P2"])

        self.assertIn("B2", late_formula)
        self.assertIn("H2", late_formula)
        self.assertIn('"[hh]:mm"', late_formula)
        self.assertIn("C2", early_formula)
        self.assertIn("I2", early_formula)
        self.assertIn('"[hh]:mm"', early_formula)
        self.assertIn("B2", work_formula)
        self.assertIn("C2", work_formula)
        self.assertIn('"[hh]:mm"', work_formula)
        self.assertIn("TEXTSPLIT(E2", work_formula)
        self.assertIn("TEXTSPLIT(D2", work_formula)
        self.assertIn("B2", basic_formula)
        self.assertIn("C2", basic_formula)
        self.assertIn("H2", basic_formula)
        self.assertIn("I2", basic_formula)
        self.assertIn("TIME(1,0,0)", basic_formula)
        self.assertIn('"[hh]:mm"', basic_formula)
        self.assertIn("TEXTSPLIT(E2", break_formula)
        self.assertIn("TEXTSPLIT(D2", break_formula)
        self.assertIn('"[hh]:mm"', break_formula)
        self.assertIn("TEXTSPLIT(G2", lunch_formula)
        self.assertIn("TEXTSPLIT(F2", lunch_formula)
        self.assertIn('"[hh]:mm"', lunch_formula)
        self.assertIn("C2", overtime_formula)
        self.assertIn("I2", overtime_formula)
        self.assertIn('"[hh]:mm"', overtime_formula)
        for formula in (
            late_formula,
            early_formula,
            work_formula,
            basic_formula,
            break_formula,
            lunch_formula,
            overtime_formula,
        ):
            self.assertNotIn("@", formula)

        for cell_ref in ("J3", "K3", "L3", "M3", "N3", "O3", "P3"):
            self.assertIn(sheet[cell_ref].value, (None, ""))

    def test_work_basic_and_overtime_formulas_zero_rest_day_and_holiday_rows(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_work_hours", "Work Hours"),
            ("daily_basic_hours", "Basic Hours"),
            ("daily_overtime", "Overtime"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E002", "E003"],
                "Clock In": ["09:00", "09:00", "09:00"],
                "Clock Out": ["18:00", "18:00", ""],
                "Shift Start": ["08:00", "08:00", "08:00"],
                "Shift End": ["17:00", "17:00", "17:00"],
                "Work Hours": ["", "", ""],
                "Basic Hours": ["", "", ""],
                "Overtime": ["", "", ""],
                "__Rest Day": ["Yes", "", ""],
                "__Holiday Type": ["", "Special Holiday", "Regular Holiday"],
            }
        )

        workbook = self._formula_workbook(df, selected_columns)
        sheet = workbook["Sheet1"]

        rest_day_formulas = [
            self._cell_formula_text(sheet[cell_ref])
            for cell_ref in ("F2", "G2", "H2")
        ]

        for formula in rest_day_formulas:
            self.assertIn('TRIM(I2&"")', formula)
            self.assertIn('B2<>""', formula)
            self.assertIn('C2<>""', formula)
            self.assertIn('"00:00"', formula)
            self.assertNotIn("@", formula)

        holiday_formulas = [
            self._cell_formula_text(sheet[cell_ref])
            for cell_ref in ("F3", "G3", "H3")
        ]
        for formula in holiday_formulas:
            self.assertIn('TRIM(J3&"")', formula)
            self.assertIn('SEARCH("SPECIAL"', formula)
            self.assertIn('SEARCH("REGULAR"', formula)
            self.assertIn('B3<>""', formula)
            self.assertIn('C3<>""', formula)
            self.assertIn('"00:00"', formula)
            self.assertNotIn("@", formula)

        blank_clock_formulas = [
            self._cell_formula_text(sheet[cell_ref])
            for cell_ref in ("F4", "G4", "H4")
        ]
        for formula in blank_clock_formulas:
            self.assertIn('TRIM(J4&"")', formula)
            self.assertIn('B4<>""', formula)
            self.assertIn('C4<>""', formula)
            self.assertIn('C4=""', formula)

    def test_premium_export_columns_are_excel_formulas(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_break_in", "Break In"),
            ("daily_break_out", "Break Out"),
            ("daily_lunch_in", "Lunch In"),
            ("daily_lunch_out", "Lunch Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_shift_day", "Shift Day"),
            ("daily_holiday_type", "Holiday Type"),
            ("daily_worked_special_holiday", "Worked on Special Holiday"),
            ("daily_ot_worked_special_holiday", "OT on Worked on Special Holiday"),
            ("daily_worked_regular_holiday", "Worked on Regular Holiday"),
            ("daily_ot_worked_regular_holiday", "OT on Worked on Regular Holiday"),
            ("daily_worked_rest_day", "Worked on Restday"),
            ("daily_ot_worked_rest_day", "OT on Worked Restday"),
            ("daily_night_differential", "Night Differential Hours"),
            (
                "daily_night_differential_overtime",
                "Night Differential Hours- OVERTIME",
            ),
            (
                "daily_night_differential_rest_day_overtime",
                "Night Differential - Rest Day Overtime",
            ),
            (
                "daily_night_differential_rest_day",
                "Night Differential Hours-REST DAY",
            ),
            (
                "daily_night_differential_regular_holiday",
                "Night Differential Regular Holiday - Hours",
            ),
            (
                "daily_night_differential_special_holiday",
                "Night Differential Special Holiday - Hours",
            ),
            (
                "daily_night_differential_special_holiday_overtime",
                "Night Differential Hours-SPECIAL HOL. OVERTIME",
            ),
            (
                "daily_night_differential_regular_holiday_overtime",
                "Night Differential - Overtime - Legal Hours",
            ),
            ("daily_rest_day_regular_holiday", "Rest Day Hours- Regular Holiday"),
            (
                "daily_ot_regular_holiday_rest_day",
                "Overtime hours - Regular Holiday - Rest Day",
            ),
            (
                "daily_night_differential_rest_day_regular_holiday_overtime",
                "Night Differential Hours-REST DAY Regular HOL. OVERTIME",
            ),
            (
                "daily_night_differential_rest_day_regular_holiday",
                "Night Differential Hours-REST DAY Regular Pay.",
            ),
            ("daily_rest_day_special_holiday", "Rest Day Hours - Special Holiday"),
            (
                "daily_night_differential_rest_day_special_holiday",
                "Night Differential Hours-REST DAY SPECIAL HOL.",
            ),
            (
                "daily_night_differential_rest_day_special_holiday_overtime",
                "Night Differential Hours-REST DAY Special HOL. OVERTIME",
            ),
            (
                "daily_ot_special_holiday_rest_day",
                "Overtime hours - Special Holiday - Rest Day",
            ),
            ("daily_leave_type", "Leave Type"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001"] * 7,
                "Clock In": ["20:00", "20:00", "20:00", "20:00", "20:00", "20:00", ""],
                "Clock Out": ["06:00", "06:00", "06:00", "06:00", "06:00", "06:00", ""],
                "Break In": [""] * 7,
                "Break Out": [""] * 7,
                "Lunch In": [""] * 7,
                "Lunch Out": [""] * 7,
                "Shift Start": ["20:00", "20:00", "20:00", "20:00", "20:00", "20:00", ""],
                "Shift End": ["04:00", "04:00", "04:00", "04:00", "04:00", "04:00", ""],
                "Shift Day": [
                    "Friday",
                    "Friday",
                    "Saturday",
                    "Saturday",
                    "Saturday",
                    "Friday",
                    "Friday",
                ],
                "Holiday Type": [
                    "Special Holiday",
                    "Regular Holiday",
                    "",
                    "Regular Holiday",
                    "Special Holiday",
                    "",
                    "",
                ],
                "Worked on Special Holiday": [""] * 7,
                "OT on Worked on Special Holiday": [""] * 7,
                "Worked on Regular Holiday": [""] * 7,
                "OT on Worked on Regular Holiday": [""] * 7,
                "Worked on Restday": [""] * 7,
                "OT on Worked Restday": [""] * 7,
                "Night Differential Hours": [""] * 7,
                "Night Differential Hours- OVERTIME": [""] * 7,
                "Night Differential - Rest Day Overtime": [""] * 7,
                "Night Differential Hours-REST DAY": [""] * 7,
                "Night Differential Regular Holiday - Hours": [""] * 7,
                "Night Differential Special Holiday - Hours": [""] * 7,
                "Night Differential Hours-SPECIAL HOL. OVERTIME": [""] * 7,
                "Night Differential - Overtime - Legal Hours": [""] * 7,
                "Rest Day Hours- Regular Holiday": [""] * 7,
                "Overtime hours - Regular Holiday - Rest Day": [""] * 7,
                "Night Differential Hours-REST DAY Regular HOL. OVERTIME": [""] * 7,
                "Night Differential Hours-REST DAY Regular Pay.": [""] * 7,
                "Rest Day Hours - Special Holiday": [""] * 7,
                "Night Differential Hours-REST DAY SPECIAL HOL.": [""] * 7,
                "Night Differential Hours-REST DAY Special HOL. OVERTIME": [""] * 7,
                "Overtime hours - Special Holiday - Rest Day": [""] * 7,
                "Leave Type": ["", "", "", "", "", "", "Vacation Leave"],
                "__Rest Day": ["", "", "Yes", "Yes", "Yes", "", ""],
            }
        )

        workbook = self._formula_workbook(df, selected_columns, [])
        sheet = workbook["Sheet1"]

        field_columns = {
            field_name: index
            for index, (field_name, _label) in enumerate(selected_columns, start=1)
        }

        def formula(field_name, row_number):
            return self._cell_formula_text(
                sheet.cell(row=row_number, column=field_columns[field_name])
            )

        special_formula = formula("daily_worked_special_holiday", 2)
        regular_formula = formula("daily_worked_regular_holiday", 3)
        rest_formula = formula("daily_worked_rest_day", 4)
        regular_rest_formula = formula("daily_rest_day_regular_holiday", 5)
        special_rest_formula = formula("daily_rest_day_special_holiday", 6)
        ordinary_nd_formula = formula("daily_night_differential", 7)
        leave_formula = formula("daily_worked_special_holiday", 8)

        self.assertNotIn("daily_rest_day", field_columns)
        self.assertNotIn("daily_shift_schedule", field_columns)
        hidden_rest_day_cell = sheet.cell(
            row=2, column=len(df.columns)
        ).coordinate
        hidden_rest_day_col = hidden_rest_day_cell.rstrip("2")

        self.assertIn('SEARCH("Special",K2)', special_formula)
        self.assertIn(f'TRIM({hidden_rest_day_col}2&"")', special_formula)
        self.assertIn('""', special_formula)
        self.assertIn('SEARCH("Regular",K3)', regular_formula)
        self.assertIn(f'TRIM({hidden_rest_day_col}4&"")', rest_formula)
        self.assertIn('SEARCH("Regular",K5)', regular_rest_formula)
        self.assertIn(f'TRIM({hidden_rest_day_col}5&"")', regular_rest_formula)
        self.assertIn('SEARCH("Special",K6)', special_rest_formula)
        self.assertIn(f'TRIM({hidden_rest_day_col}6&"")', special_rest_formula)
        self.assertIn('NOT(ISNUMBER(SEARCH("Regular",K7)))', ordinary_nd_formula)
        self.assertIn('NOT(ISNUMBER(SEARCH("Special",K7)))', ordinary_nd_formula)
        self.assertIn("TIME(22,0,0)", ordinary_nd_formula)
        for formula_text in (
            special_formula,
            regular_formula,
            rest_formula,
            regular_rest_formula,
            special_rest_formula,
            ordinary_nd_formula,
        ):
            self.assertTrue(formula_text.startswith("="))
            self.assertNotIn("@", formula_text)
        self.assertIn(leave_formula, (None, ""))

    def test_premium_export_formulas_prefer_visible_rest_day_column(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_rest_day", "Rest Day"),
            ("daily_holiday_type", "Holiday Type"),
            ("daily_worked_rest_day", "Worked on Restday"),
            ("daily_ot_worked_rest_day", "OT on Worked Restday"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001"],
                "Clock In": ["09:00"],
                "Clock Out": ["18:00"],
                "Shift Start": ["08:00"],
                "Shift End": ["17:00"],
                "Rest Day": ["Yes"],
                "Holiday Type": [""],
                "Worked on Restday": [""],
                "OT on Worked Restday": [""],
            }
        )

        workbook = self._formula_workbook(df, selected_columns, [])
        sheet = workbook["Sheet1"]

        rest_formula = self._cell_formula_text(sheet["H2"])
        overtime_formula = self._cell_formula_text(sheet["I2"])

        self.assertIn('TRIM(F2&"")', rest_formula)
        self.assertIn('TRIM(F2&"")', overtime_formula)
        self.assertNotIn("__Rest Day", [cell.value for cell in sheet[1]])

    def test_premium_export_formulas_allow_blank_shift_times_for_rest_day_work(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_holiday_type", "Holiday Type"),
            ("daily_worked_rest_day", "Worked on Restday"),
            ("daily_ot_worked_rest_day", "OT on Worked Restday"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001"],
                "Clock In": ["09:00"],
                "Clock Out": ["17:00"],
                "Shift Start": [""],
                "Shift End": [""],
                "Holiday Type": [""],
                "Worked on Restday": [""],
                "OT on Worked Restday": [""],
                "__Rest Day": ["Yes"],
            }
        )

        workbook = self._formula_workbook(df, selected_columns, [])
        sheet = workbook["Sheet1"]

        rest_formula = self._cell_formula_text(sheet["G2"])
        overtime_formula = self._cell_formula_text(sheet["H2"])

        self.assertIn('=IF(AND(AND(B2<>"",C2<>""),', rest_formula)
        self.assertIn("IF(AND(D2<>\"\",E2<>\"\"),MIN", rest_formula)
        self.assertIn('""', overtime_formula)

    def test_totals_sheet_groups_by_employee_and_points_to_employee_ranges(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("employee_id", "Employee"),
            ("daily_basic_hours", "Basic Hours"),
            ("daily_overtime", "Overtime"),
            ("daily_worked_special_holiday", "Worked on Special Holiday"),
            ("daily_late_come", "Late Come"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E001", "E002"],
                "Employee": ["Ada Lovelace", "Ada Lovelace", "Grace Hopper"],
                "Basic Hours": ["08:00", "08:00", "08:00"],
                "Overtime": ["00:30", "01:00", "00:15"],
                "Worked on Special Holiday": ["02:00", "00:30", "00:00"],
                "Late Come": ["00:05", "00:10", "00:03"],
            }
        )

        total_ranges = _build_export_total_ranges(df, selected_columns)
        workbook = self._formula_workbook(
            df, selected_columns, total_ranges, totals_sheet_name="Totals"
        )
        sheet = workbook["Sheet1"]
        totals_sheet = workbook["Totals"]
        headers = [cell.value for cell in totals_sheet[1]]

        self.assertEqual(
            sheet["A2"].value,
            "E001",
        )
        self.assertEqual(sheet["A4"].value, "E002")
        self.assertNotIn("Total", df["Employee No."].tolist())
        self.assertEqual(
            total_ranges,
            [
                {
                    "start_df_row": 0,
                    "end_df_row": 1,
                    "employee_number": "E001",
                    "employee_name": "Ada Lovelace",
                },
                {
                    "start_df_row": 2,
                    "end_df_row": 2,
                    "employee_number": "E002",
                    "employee_name": "Grace Hopper",
                },
            ],
        )

        basic_col = headers.index("Total Basic Hours") + 1
        overtime_col = headers.index("OT Hours") + 1
        special_col = headers.index("Worked on Special Holiday") + 1
        late_col = headers.index("Late") + 1

        self.assertEqual(totals_sheet["A2"].value, "E001")
        self.assertEqual(totals_sheet["B2"].value, "Ada Lovelace")
        formula_ranges = (
            (totals_sheet.cell(2, basic_col), "'Sheet1'!C2:C3"),
            (totals_sheet.cell(2, overtime_col), "'Sheet1'!D2:D3"),
            (totals_sheet.cell(2, special_col), "'Sheet1'!E2:E3"),
            (totals_sheet.cell(2, late_col), "'Sheet1'!F2:F3"),
            (totals_sheet.cell(3, basic_col), "'Sheet1'!C4:C4"),
        )
        for cell, expected_range in formula_ranges:
            formula = self._cell_formula_text(cell)
            self.assertTrue(formula.startswith("=SUMPRODUCT("))
            self.assertIn(f"IFERROR(N({expected_range}),0)", formula)
            self.assertIn(f"LEFT({expected_range},FIND", formula)
            self.assertIn(f"MID({expected_range},FIND", formula)
            self.assertNotIn("TIMEVALUE", formula)
            self.assertNotIn("@", formula)

    def test_employee_separator_rows_shift_totals_and_keep_formulas_off_blank_rows(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("employee_id", "Employee"),
            ("daily_clock_in", "Clock In"),
            ("daily_clock_out", "Clock Out"),
            ("daily_shift_start", "Shift Start"),
            ("daily_shift_end", "Shift End"),
            ("daily_basic_hours", "Basic Hours"),
            ("daily_overtime", "Overtime"),
            ("daily_late_come", "Late Come"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E001", "E002"],
                "Employee": ["Ada Lovelace", "Ada Lovelace", "Grace Hopper"],
                "Clock In": ["09:00", "09:00", "08:00"],
                "Clock Out": ["18:00", "18:30", "17:15"],
                "Shift Start": ["08:00", "08:00", "08:00"],
                "Shift End": ["17:00", "17:00", "17:00"],
                "Basic Hours": ["", "", ""],
                "Overtime": ["", "", ""],
                "Late Come": ["", "", ""],
            }
        )

        total_ranges = _build_export_total_ranges(df, selected_columns)
        spaced_df, total_ranges = _insert_export_employee_separator_rows(
            df, total_ranges
        )
        workbook = self._formula_workbook(
            spaced_df, selected_columns, total_ranges, totals_sheet_name="Totals"
        )
        sheet = workbook["Sheet1"]
        totals_sheet = workbook["Totals"]
        headers = [cell.value for cell in totals_sheet[1]]

        self.assertEqual(
            spaced_df["Employee No."].tolist(),
            ["E001", "E001", "", "E002"],
        )
        self.assertEqual(
            total_ranges,
            [
                {
                    "start_df_row": 0,
                    "end_df_row": 1,
                    "employee_number": "E001",
                    "employee_name": "Ada Lovelace",
                },
                {
                    "start_df_row": 3,
                    "end_df_row": 3,
                    "employee_number": "E002",
                    "employee_name": "Grace Hopper",
                },
            ],
        )
        self.assertEqual(sheet["A2"].value, "E001")
        self.assertEqual(sheet["A4"].value, None)
        self.assertEqual(sheet["A5"].value, "E002")
        for cell_ref in ("G4", "H4", "I4"):
            self.assertIn(sheet[cell_ref].value, (None, ""))

        basic_col = headers.index("Total Basic Hours") + 1
        overtime_col = headers.index("OT Hours") + 1
        late_col = headers.index("Late") + 1
        expected_ranges = (
            (totals_sheet.cell(2, basic_col), "'Sheet1'!G2:G3"),
            (totals_sheet.cell(3, basic_col), "'Sheet1'!G5:G5"),
            (totals_sheet.cell(3, overtime_col), "'Sheet1'!H5:H5"),
            (totals_sheet.cell(3, late_col), "'Sheet1'!I5:I5"),
        )
        for cell, expected_range in expected_ranges:
            formula = self._cell_formula_text(cell)
            self.assertIn(f"IFERROR(N({expected_range}),0)", formula)
            self.assertNotIn("'Sheet1'!G4:G4", formula)
            self.assertNotIn("@", formula)

    def test_employee_separator_rows_do_not_add_trailing_blank_for_one_employee(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("employee_id", "Employee"),
            ("daily_basic_hours", "Basic Hours"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E001"],
                "Employee": ["Ada Lovelace", "Ada Lovelace"],
                "Basic Hours": ["08:00", "08:00"],
            }
        )

        total_ranges = _build_export_total_ranges(df, selected_columns)
        spaced_df, spaced_ranges = _insert_export_employee_separator_rows(
            df, total_ranges
        )

        self.assertEqual(spaced_df["Employee No."].tolist(), ["E001", "E001"])
        self.assertEqual(spaced_ranges, total_ranges)

    def test_totals_sheet_leaves_missing_total_sources_blank(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("attendance_date", "Attendance Date"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": ["E001", "E002"],
                "Attendance Date": [date(2026, 4, 10), date(2026, 4, 11)],
            }
        )

        total_ranges = _build_export_total_ranges(df, selected_columns)
        workbook = self._formula_workbook(
            df, selected_columns, total_ranges, totals_sheet_name="Totals"
        )
        sheet = workbook["Sheet1"]
        totals_sheet = workbook["Totals"]

        self.assertEqual(
            [sheet["A2"].value, sheet["A3"].value],
            ["E001", "E002"],
        )
        self.assertEqual(totals_sheet["A2"].value, "E001")
        self.assertIsNone(totals_sheet["C2"].value)
        self.assertIsNone(totals_sheet["C3"].value)

    def test_empty_totals_sheet_has_headers_only(self):
        selected_columns = [
            ("employee_number", "Employee No."),
            ("daily_late_come", "Late Come"),
            ("daily_basic_hours", "Basic Hours"),
            ("daily_worked_special_holiday", "Worked on Special Holiday"),
        ]
        df = pd.DataFrame(
            {
                "Employee No.": [],
                "Late Come": [],
                "Basic Hours": [],
                "Worked on Special Holiday": [],
            }
        )

        total_ranges = _build_export_total_ranges(df, selected_columns)
        workbook = self._formula_workbook(
            df, selected_columns, total_ranges, totals_sheet_name="Totals"
        )
        totals_sheet = workbook["Totals"]

        self.assertEqual(total_ranges, [])
        self.assertEqual(totals_sheet["A1"].value, "EMP No.")
        self.assertEqual(totals_sheet["C1"].value, "Total Basic Hours")
        self.assertEqual(totals_sheet.max_row, 1)


class DailyActivityRowsTests(SimpleTestCase):
    databases = {"default"}

    def _template_employee(self):
        return SimpleNamespace(
            id=101,
            get_avatar="/media/avatar.jpg",
            employee_work_info=SimpleNamespace(
                branch_id="Main Branch",
                department_id="HR",
                job_position_id="Analyst",
            ),
        )

    def _image(self, url):
        return SimpleNamespace(url=url)

    def _activity(
        self,
        pk,
        employee,
        activity_type,
        clock_in,
        clock_out=None,
        clock_in_selfie=None,
        clock_out_selfie=None,
        shift_day="friday",
        in_datetime=None,
        out_datetime=None,
        attendance_date=None,
        clock_out_date=None,
    ):
        attendance_date = attendance_date or date(2026, 4, 10)
        return SimpleNamespace(
            id=pk,
            employee_id=employee,
            employee_id_id=employee.id,
            attendance_date=attendance_date,
            shift_day=shift_day,
            activity_type=activity_type,
            clock_in_date=attendance_date,
            clock_in=clock_in,
            in_datetime=in_datetime or datetime.combine(attendance_date, clock_in),
            clock_out_date=clock_out_date or (attendance_date if clock_out else None),
            clock_out=clock_out,
            out_datetime=out_datetime
            or (datetime.combine(attendance_date, clock_out) if clock_out else None),
            clock_in_selfie=clock_in_selfie,
            clock_out_selfie=clock_out_selfie,
            clock_in_gps_address="Clock In Address",
            clock_out_gps_address="Clock Out Address" if clock_out else None,
            gps_address=None,
            latitude=None,
            longitude=None,
            clock_in_maps_url="https://maps.example/in",
            clock_out_maps_url="https://maps.example/out" if clock_out else None,
        )

    def _multi_work_image_row(self):
        employee = self._template_employee()
        with patch("attendance.views.views.Attendance") as attendance_model:
            attendance_model.objects.filter.return_value = []
            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            1,
                            employee,
                            "work",
                            time(9, 0),
                            time(12, 0),
                            self._image("/media/work-1-in.jpg"),
                            self._image("/media/work-1-out.jpg"),
                        ),
                        self._activity(
                            2,
                            employee,
                            "work",
                            time(13, 0),
                            time(17, 0),
                            self._image("/media/work-2-in.jpg"),
                            self._image("/media/work-2-out.jpg"),
                        ),
                    ]
                )
            )
        return rows[0]

    @patch("attendance.views.views.Attendance")
    def test_build_daily_rows_groups_work_breaks_lunch_and_images(self, attendance_model):
        employee = SimpleNamespace(id=101)
        attendance = SimpleNamespace(
            id=None,
            employee_id_id=101,
            attendance_date=date(2026, 4, 10),
            shift_id="Morning",
            work_type_id="Office",
            minimum_hour="08:00",
            attendance_worked_hour="07:30",
            attendance_overtime="00:00",
            hours_pending=lambda: "00:30",
        )
        attendance_model.objects.filter.return_value = [attendance]
        break_in_image = SimpleNamespace(url="/media/break-in.jpg")
        break_out_image = SimpleNamespace(url="/media/break-out.jpg")

        rows = build_daily_activity_rows(
            FakeActivityQuerySet(
                [
                    self._activity(1, employee, "work", datetime(2026, 4, 10, 9).time(), datetime(2026, 4, 10, 17).time()),
                    self._activity(2, employee, "break", datetime(2026, 4, 10, 10).time(), datetime(2026, 4, 10, 10, 15).time(), break_in_image, break_out_image),
                    self._activity(3, employee, "break", datetime(2026, 4, 10, 15).time(), datetime(2026, 4, 10, 15, 10).time()),
                    self._activity(4, employee, "lunch", datetime(2026, 4, 10, 12).time(), datetime(2026, 4, 10, 13).time()),
                ]
            )
        )

        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.activity_ids, [1, 2, 3, 4])
        self.assertEqual(row.detail_activity_id, 1)
        self.assertEqual(row.work_in.activity.id, 1)
        self.assertEqual(row.work_out.activity.id, 1)
        self.assertEqual(len(row.break_segments), 2)
        self.assertEqual(len(row.lunch_segments), 1)
        self.assertEqual(row.break_segments[0].clock_in_selfie, break_in_image)
        self.assertEqual(row.break_segments[0].clock_out_selfie, break_out_image)
        self.assertEqual(row.shift, "Morning")
        self.assertEqual(row.work_hours, "07:30")
        self.assertEqual(row.pending_hour, "00:30")
        self.assertEqual(row.late_come_duration, "")
        self.assertEqual(row.early_out_duration, "")

    def test_build_daily_rows_treats_blank_shift_schedule_as_rest_day(self):
        saturday = date(2026, 4, 11)
        shift_day = SimpleNamespace(id=6, day="saturday")
        employee = SimpleNamespace(
            id=118,
            employee_work_info=SimpleNamespace(shift_id="Morning", shift_id_id=1),
        )
        weekday_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=1,
            day=SimpleNamespace(day="monday"),
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        blank_rest_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=6,
            day=SimpleNamespace(day="saturday"),
            start_time=None,
            end_time=None,
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = []
            schedule_model.objects.filter.return_value = [
                weekday_schedule,
                blank_rest_schedule,
            ]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            42,
                            employee,
                            "work",
                            time(9, 0),
                            time(18, 0),
                            shift_day=shift_day,
                            attendance_date=saturday,
                        )
                    ]
                )
            )

        self.assertTrue(rows[0].is_rest_day)
        self.assertEqual(rows[0].shift_schedule, "Mon")
        self.assertEqual(rows[0].overtime, "01:00")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_start", None
                ),
                "08:00",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_end", None
                ),
                "17:00",
            )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_worked_rest_day", None
            ),
            "08:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_ot_worked_rest_day", None
            ),
            "01:00",
        )
        if hasattr(rows[0], "_premium_export_buckets"):
            delattr(rows[0], "_premium_export_buckets")
        rows[0].holiday_type = "regular"
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_rest_day_regular_holiday", None
            ),
            "08:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_ot_regular_holiday_rest_day", None
            ),
            "01:00",
        )
        if hasattr(rows[0], "_premium_export_buckets"):
            delattr(rows[0], "_premium_export_buckets")
        rows[0].holiday_type = "special"
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_rest_day_special_holiday", None
            ),
            "08:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_ot_special_holiday_rest_day", None
            ),
            "01:00",
        )

    def test_build_daily_rows_treats_explicit_schedule_flag_as_rest_day(self):
        saturday = date(2026, 4, 11)
        shift_day = SimpleNamespace(id=6, day="saturday")
        employee = SimpleNamespace(
            id=119,
            employee_work_info=SimpleNamespace(shift_id="Morning", shift_id_id=1),
        )
        weekday_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=1,
            day=SimpleNamespace(day="monday"),
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
            is_rest_day=False,
        )
        explicit_rest_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=6,
            day=SimpleNamespace(day="saturday"),
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
            is_rest_day=True,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = []
            schedule_model.objects.filter.return_value = [
                weekday_schedule,
                explicit_rest_schedule,
            ]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            43,
                            employee,
                            "work",
                            time(9, 0),
                            time(18, 0),
                            shift_day=shift_day,
                            attendance_date=saturday,
                        )
                    ]
                )
            )

        self.assertTrue(rows[0].is_rest_day)
        self.assertTrue(rows[0].has_shift_schedule)
        self.assertEqual(rows[0].shift_schedule, "Mon")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_start", None
                ),
                "08:00",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_end", None
                ),
                "17:00",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_rest_day", None
                ),
                "Yes",
            )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_worked_rest_day", None
            ),
            "08:00",
        )
        self.assertEqual(
            _attendance_activity_daily_export_value(
                rows[0], "daily_ot_worked_rest_day", None
            ),
            "01:00",
        )

    def test_build_daily_rows_keeps_all_work_segment_images(self):
        row = self._multi_work_image_row()

        self.assertEqual(len(row.work_segments), 2)
        self.assertTrue(row.has_work_images)
        self.assertEqual(row.work_segments[0].clock_in_selfie.url, "/media/work-1-in.jpg")
        self.assertEqual(row.work_segments[0].clock_out_selfie.url, "/media/work-1-out.jpg")
        self.assertEqual(row.work_segments[1].clock_in_selfie.url, "/media/work-2-in.jpg")
        self.assertEqual(row.work_segments[1].clock_out_selfie.url, "/media/work-2-out.jpg")

    def test_daily_activity_row_renders_all_work_segment_images(self):
        row = self._multi_work_image_row()
        html = render_to_string(
            "attendance/attendance_activity/daily_activity_row.html",
            {"row": row, "pd": ""},
        )

        for image_url in (
            "/media/work-1-in.jpg",
            "/media/work-1-out.jpg",
            "/media/work-2-in.jpg",
            "/media/work-2-out.jpg",
        ):
            self.assertIn(image_url, html)

        clock_image_cell = html.split('data-cell-index="6"', 1)[1].split(
            'data-cell-index="7"', 1
        )[0]
        self.assertNotIn("&mdash;", clock_image_cell)

    def test_daily_activity_modal_renders_all_work_segment_images(self):
        row = self._multi_work_image_row()
        html = render_to_string(
            "attendance/attendance_activity/daily_attendance_activity.html",
            {"row": row, "pd": ""},
        )

        for image_url in (
            "/media/work-1-in.jpg",
            "/media/work-1-out.jpg",
            "/media/work-2-in.jpg",
            "/media/work-2-out.jpg",
        ):
            self.assertIn(image_url, html)

    @patch("attendance.views.views.Attendance")
    def test_build_daily_rows_handles_missing_attendance_and_open_break(self, attendance_model):
        employee = SimpleNamespace(id=102)
        attendance_model.objects.filter.return_value = []

        rows = build_daily_activity_rows(
            FakeActivityQuerySet(
                [
                    self._activity(
                        10,
                        employee,
                        "break",
                        datetime(2026, 4, 10, 10).time(),
                    )
                ]
            )
        )

        row = rows[0]
        self.assertEqual(row.activity_ids, [10])
        self.assertEqual(row.detail_activity_id, 10)
        self.assertEqual(len(row.break_segments), 1)
        self.assertIsNone(row.break_segments[0].clock_out)
        self.assertEqual(row.work_hours, "")
        self.assertIsNone(row.shift)
        self.assertEqual(row.late_come_duration, "")
        self.assertEqual(row.early_out_duration, "")

    def test_build_daily_rows_calculates_late_and_early_durations_without_reports(self):
        employee = SimpleNamespace(id=103)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=201,
            employee_id_id=103,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(15, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(15, 0),
            hours_pending=lambda: "06:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            20,
                            employee,
                            "work",
                            time(15, 0),
                            time(15, 0),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "07:00")
        self.assertEqual(rows[0].early_out_duration, "02:00")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_late_come", None
                ),
                "07:00",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_early_out", None
                ),
                "02:00",
            )
        self.assertEqual(rows[0].detail_activity_id, 20)

    def test_build_daily_rows_uses_weekday_schedule_fallback_when_exact_day_missing(self):
        employee = SimpleNamespace(id=118)
        attendance_date = date(2026, 4, 11)
        attendance = SimpleNamespace(
            id=214,
            employee_id_id=118,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=6,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(10, 30),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(15, 0),
            hours_pending=lambda: "06:00",
        )
        weekday_schedules = [
            SimpleNamespace(
                shift_id_id=1,
                day_id=day_id,
                day=SimpleNamespace(day=day_name),
                start_time=time(8, 0),
                end_time=time(17, 0),
                is_night_shift=False,
            )
            for day_id, day_name in enumerate(
                ["monday", "tuesday", "wednesday", "thursday", "friday"],
                start=1,
            )
        ]

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = weekday_schedules

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            42,
                            employee,
                            "work",
                            time(10, 30),
                            time(15, 0),
                            shift_day="saturday",
                            attendance_date=attendance_date,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "02:30")
        self.assertEqual(rows[0].early_out_duration, "02:00")
        self.assertTrue(rows[0].is_rest_day)
        self.assertFalse(rows[0].has_shift_schedule)
        self.assertEqual(rows[0].shift_schedule, "Mon, Tue, Wed, Thu, Fri")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_start", None
                ),
                "",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_end", None
                ),
                "",
            )

    def test_build_daily_rows_prefers_exact_day_schedule_over_weekday_fallback(self):
        employee = SimpleNamespace(id=119)
        attendance_date = date(2026, 4, 11)
        attendance = SimpleNamespace(
            id=215,
            employee_id_id=119,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=6,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(11, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(16, 0),
            hours_pending=lambda: "06:00",
        )
        monday_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=1,
            day=SimpleNamespace(day="monday"),
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        saturday_schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=6,
            day=SimpleNamespace(day="saturday"),
            start_time=time(10, 0),
            end_time=time(19, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [
                monday_schedule,
                saturday_schedule,
            ]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            43,
                            employee,
                            "work",
                            time(11, 0),
                            time(16, 0),
                            shift_day="saturday",
                            attendance_date=attendance_date,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "01:00")
        self.assertEqual(rows[0].early_out_duration, "03:00")
        self.assertFalse(rows[0].is_rest_day)
        self.assertTrue(rows[0].has_shift_schedule)
        self.assertEqual(rows[0].shift_schedule, "Mon, Sat")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_start", None
                ),
                "10:00",
            )
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_shift_end", None
                ),
                "19:00",
            )

    def test_build_daily_rows_uses_activity_clock_in_when_attendance_is_stale(self):
        employee = SimpleNamespace(id=106)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=204,
            employee_id_id=106,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "06:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            23,
                            employee,
                            "work",
                            time(15, 0),
                            time(15, 0),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "07:00")
        self.assertEqual(rows[0].early_out_duration, "02:00")

    def test_build_daily_rows_uses_first_work_clock_in_for_late_come(self):
        employee = SimpleNamespace(id=110)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=207,
            employee_id_id=110,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "06:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(27, employee, "break", time(15, 0), time(15, 0)),
                        self._activity(28, employee, "work", time(16, 0), time(17, 0)),
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "08:00")

    def test_build_daily_rows_uses_last_work_clock_out_for_early_out(self):
        employee = SimpleNamespace(id=111)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=208,
            employee_id_id=111,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="07:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "01:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(29, employee, "work", time(8, 0), time(14, 0)),
                        self._activity(30, employee, "lunch", time(14, 30), time(15, 0)),
                    ]
                )
            )

        self.assertEqual(rows[0].early_out_duration, "03:00")

    def test_build_daily_rows_uses_split_work_segments_for_late_and_early(self):
        employee = SimpleNamespace(id=113)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=209,
            employee_id_id=113,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="06:30",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(9, 30),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "01:30",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(9, 0),
            end_time=time(18, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(32, employee, "work", time(9, 30), time(12, 0)),
                        self._activity(33, employee, "lunch", time(12, 0), time(13, 0)),
                        self._activity(34, employee, "work", time(13, 0), time(17, 0)),
                    ]
                )
            )

        self.assertEqual(rows[0].work_in.activity.id, 32)
        self.assertEqual(rows[0].work_out.activity.id, 34)
        self.assertEqual(rows[0].late_come_duration, "00:30")
        self.assertEqual(rows[0].early_out_duration, "01:00")

    def test_build_daily_rows_keeps_early_out_blank_when_resumed_work_is_open(self):
        employee = SimpleNamespace(id=114)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=210,
            employee_id_id=114,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="02:30",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(9, 30),
            attendance_clock_out_date=None,
            attendance_clock_out=None,
            hours_pending=lambda: "05:30",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(9, 0),
            end_time=time(18, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(35, employee, "work", time(9, 30), time(12, 0)),
                        self._activity(36, employee, "break", time(12, 0), time(12, 15)),
                        self._activity(37, employee, "work", time(12, 15)),
                    ]
                )
            )

        self.assertEqual(rows[0].work_in.activity.id, 35)
        self.assertEqual(rows[0].work_out.activity.id, 35)
        self.assertEqual(rows[0].late_come_duration, "00:30")
        self.assertEqual(rows[0].early_out_duration, "")

    def test_build_daily_rows_ignores_later_break_or_lunch_out_for_early_out(self):
        employee = SimpleNamespace(id=115)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=211,
            employee_id_id=115,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="03:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(9, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(12, 0),
            hours_pending=lambda: "05:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(9, 0),
            end_time=time(18, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(38, employee, "work", time(9, 0), time(12, 0)),
                        self._activity(39, employee, "lunch", time(12, 0), time(13, 0)),
                    ]
                )
            )

        self.assertEqual(rows[0].work_out.activity.id, 38)
        self.assertEqual(rows[0].early_out_duration, "06:00")

    def test_build_daily_rows_handles_aware_clock_in_datetime(self):
        employee = SimpleNamespace(id=108)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=205,
            employee_id_id=108,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="02:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(15, 0),
            hours_pending=lambda: "06:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        aware_clock_in = timezone.make_aware(datetime(2026, 4, 10, 15, 0))

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            25,
                            employee,
                            "work",
                            time(15, 0),
                            time(15, 0),
                            in_datetime=aware_clock_in,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "07:00")

    def test_build_daily_rows_handles_aware_clock_out_datetime(self):
        employee = SimpleNamespace(id=109)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=206,
            employee_id_id=109,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="07:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(15, 0),
            hours_pending=lambda: "01:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        aware_clock_out = timezone.make_aware(datetime(2026, 4, 10, 15, 0))

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            26,
                            employee,
                            "work",
                            time(8, 0),
                            time(15, 0),
                            out_datetime=aware_clock_out,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "")
        self.assertEqual(rows[0].early_out_duration, "02:00")

    def test_build_daily_rows_late_come_uses_displayed_clock_in_time(self):
        employee = SimpleNamespace(id=116)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=212,
            employee_id_id=116,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="00:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(14, 56),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "08:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            40,
                            employee,
                            "work",
                            time(14, 56),
                            time(17, 0),
                            in_datetime=datetime(2026, 4, 10, 14, 13),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].work_in.clock_in, time(14, 56))
        self.assertEqual(rows[0].late_come_duration, "06:56")

    def test_build_daily_rows_early_out_uses_displayed_clock_out_time(self):
        employee = SimpleNamespace(id=117)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=213,
            employee_id_id=117,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="00:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(14, 56),
            hours_pending=lambda: "08:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            41,
                            employee,
                            "work",
                            time(8, 0),
                            time(14, 56),
                            out_datetime=datetime(2026, 4, 10, 14, 13),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].work_out.clock_out, time(14, 56))
        self.assertEqual(rows[0].early_out_duration, "02:04")

    def test_build_daily_rows_calculates_late_with_employee_shift_fallback(self):
        shift_day = SimpleNamespace(id=5)
        employee = SimpleNamespace(
            id=107,
            employee_work_info=SimpleNamespace(shift_id="Morning", shift_id_id=1),
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = []
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            24,
                            employee,
                            "work",
                            time(15, 0),
                            time(15, 0),
                            shift_day=shift_day,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].shift, "Morning")
        self.assertEqual(rows[0].shift_day, shift_day)
        self.assertEqual(rows[0].late_come_duration, "07:00")
        self.assertEqual(rows[0].early_out_duration, "02:00")

    def test_build_daily_rows_uses_attendance_date_weekday_when_shift_day_missing(self):
        employee = SimpleNamespace(
            id=112,
            employee_work_info=SimpleNamespace(shift_id="Morning", shift_id_id=1),
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            day=SimpleNamespace(day="friday"),
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = []
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            31,
                            employee,
                            "work",
                            time(15, 0),
                            time(15, 0),
                            shift_day=None,
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].shift, "Morning")
        self.assertEqual(rows[0].shift_day, "friday")
        self.assertEqual(rows[0].late_come_duration, "07:00")
        self.assertEqual(rows[0].early_out_duration, "02:00")

    def test_build_daily_rows_keeps_empty_late_early_durations_when_on_time(self):
        employee = SimpleNamespace(id=104)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=202,
            employee_id_id=104,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="08:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(9, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(17, 0),
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(9, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [self._activity(21, employee, "work", time(9, 0), time(17, 0))]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "")
        self.assertEqual(rows[0].early_out_duration, "")

    def test_build_daily_rows_uses_next_day_for_night_shift_early_out(self):
        employee = SimpleNamespace(id=105)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=203,
            employee_id_id=105,
            attendance_date=attendance_date,
            shift_id="Night",
            shift_id_id=2,
            work_type_id="Office",
            attendance_day_id=6,
            minimum_hour="08:00",
            attendance_worked_hour="07:30",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(22, 0),
            attendance_clock_out_date=attendance_date + timedelta(days=1),
            attendance_clock_out=time(5, 30),
            hours_pending=lambda: "00:30",
        )
        schedule = SimpleNamespace(
            shift_id_id=2,
            day_id=6,
            start_time=time(22, 0),
            end_time=time(6, 0),
            is_night_shift=True,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            22,
                            employee,
                            "work",
                            time(22, 0),
                            time(5, 30),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].late_come_duration, "")
        self.assertEqual(rows[0].early_out_duration, "00:30")

    def test_build_daily_rows_uses_live_schedule_end_overtime(self):
        employee = SimpleNamespace(id=120)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=216,
            employee_id_id=120,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="13:25",
            attendance_overtime="05:38",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(21, 25),
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            44,
                            employee,
                            "work",
                            time(8, 0),
                            time(21, 25),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].overtime, "04:25")
        with patch(
            "attendance.views.views.format_export_value",
            side_effect=lambda value, employee: value,
        ):
            self.assertEqual(
                _attendance_activity_daily_export_value(
                    rows[0], "daily_overtime", None
                ),
                "04:25",
            )

    def test_build_daily_rows_uses_last_work_clock_out_after_shift_breaks(self):
        employee = SimpleNamespace(id=121)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=217,
            employee_id_id=121,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="09:30",
            attendance_overtime="02:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(19, 0),
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(45, employee, "work", time(8, 0), time(17, 0)),
                        self._activity(46, employee, "break", time(17, 0), time(17, 30)),
                        self._activity(47, employee, "work", time(17, 30), time(19, 0)),
                    ]
                )
            )

        self.assertEqual(rows[0].overtime, "02:00")

    def test_build_daily_rows_ignores_open_work_for_overtime_until_clock_out(self):
        employee = SimpleNamespace(id=122)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=218,
            employee_id_id=122,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="08:00",
            attendance_overtime="00:00",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(16, 0),
            attendance_clock_out_date=None,
            attendance_clock_out=None,
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [self._activity(48, employee, "work", time(16, 0))]
                )
            )

        self.assertEqual(rows[0].overtime, "00:00")

    def test_build_daily_rows_uses_last_clock_out_when_later_work_is_open(self):
        employee = SimpleNamespace(id=124)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=220,
            employee_id_id=124,
            attendance_date=attendance_date,
            shift_id="Morning",
            shift_id_id=1,
            work_type_id="Office",
            attendance_day_id=5,
            minimum_hour="08:00",
            attendance_worked_hour="13:25",
            attendance_overtime="25:02",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(8, 0),
            attendance_clock_out_date=attendance_date,
            attendance_clock_out=time(21, 25),
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=1,
            day_id=5,
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model, patch(
            "attendance.methods.utils.django_timezone.now",
            return_value=datetime(2026, 4, 11, 18, 2),
        ):
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(50, employee, "work", time(8, 0), time(21, 25)),
                        self._activity(51, employee, "work", time(21, 30)),
                    ]
                )
            )

        self.assertEqual(rows[0].overtime, "04:25")

    def test_build_daily_rows_calculates_night_shift_overtime_after_next_day_end(self):
        employee = SimpleNamespace(id=123)
        attendance_date = date(2026, 4, 10)
        attendance = SimpleNamespace(
            id=219,
            employee_id_id=123,
            attendance_date=attendance_date,
            shift_id="Night",
            shift_id_id=2,
            work_type_id="Office",
            attendance_day_id=6,
            minimum_hour="08:00",
            attendance_worked_hour="10:30",
            attendance_overtime="02:30",
            attendance_clock_in_date=attendance_date,
            attendance_clock_in=time(22, 0),
            attendance_clock_out_date=attendance_date + timedelta(days=1),
            attendance_clock_out=time(8, 30),
            hours_pending=lambda: "00:00",
        )
        schedule = SimpleNamespace(
            shift_id_id=2,
            day_id=6,
            start_time=time(22, 0),
            end_time=time(6, 0),
            is_night_shift=True,
        )

        with patch("attendance.views.views.Attendance") as attendance_model, patch(
            "attendance.views.views.EmployeeShiftSchedule"
        ) as schedule_model:
            attendance_model.objects.filter.return_value = [attendance]
            schedule_model.objects.filter.return_value = [schedule]

            rows = build_daily_activity_rows(
                FakeActivityQuerySet(
                    [
                        self._activity(
                            49,
                            employee,
                            "work",
                            time(22, 0),
                            time(8, 30),
                            clock_out_date=attendance_date + timedelta(days=1),
                            out_datetime=datetime(2026, 4, 11, 8, 30),
                        )
                    ]
                )
            )

        self.assertEqual(rows[0].overtime, "02:30")


class AttendanceLateComeEarlyOutDurationTests(SimpleTestCase):
    def _attendance(self, attendance_day):
        return Attendance(
            attendance_date=date(2026, 4, 11),
            shift_id=EmployeeShift(id=1, employee_shift="Morning"),
            attendance_day=attendance_day,
            attendance_clock_in_date=date(2026, 4, 11),
            attendance_clock_in=time(10, 30),
            attendance_clock_out_date=date(2026, 4, 11),
            attendance_clock_out=time(15, 0),
        )

    @patch("attendance.methods.utils.EmployeeShiftSchedule")
    def test_get_late_early_duration_uses_weekday_fallback_when_exact_day_missing(
        self, schedule_model
    ):
        saturday = EmployeeShiftDay(id=6, day="saturday")
        attendance = self._attendance(saturday)
        exact_qs = MagicMock()
        exact_qs.first.return_value = None
        fallback_qs = MagicMock()
        fallback_qs.first.return_value = SimpleNamespace(
            start_time=time(8, 0),
            end_time=time(17, 0),
            is_night_shift=False,
        )
        schedule_model.objects.filter.side_effect = [
            exact_qs,
            fallback_qs,
            exact_qs,
            fallback_qs,
        ]

        late_report = AttendanceLateComeEarlyOut(
            attendance_id=attendance, type="late_come"
        )
        early_report = AttendanceLateComeEarlyOut(
            attendance_id=attendance, type="early_out"
        )

        self.assertEqual(late_report.get_late_early_duration(), "02:30")
        self.assertEqual(early_report.get_late_early_duration(), "02:00")

    @patch("attendance.methods.utils.EmployeeShiftSchedule")
    def test_get_late_early_duration_prefers_exact_day_schedule(self, schedule_model):
        saturday = EmployeeShiftDay(id=6, day="saturday")
        attendance = self._attendance(saturday)
        exact_qs = MagicMock()
        exact_qs.first.return_value = SimpleNamespace(
            start_time=time(10, 0),
            end_time=time(19, 0),
            is_night_shift=False,
        )
        schedule_model.objects.filter.return_value = exact_qs

        late_report = AttendanceLateComeEarlyOut(
            attendance_id=attendance, type="late_come"
        )
        early_report = AttendanceLateComeEarlyOut(
            attendance_id=attendance, type="early_out"
        )

        self.assertEqual(late_report.get_late_early_duration(), "00:30")
        self.assertEqual(early_report.get_late_early_duration(), "04:00")


class AttendanceExportFormFieldTests(SimpleTestCase):
    def test_attendance_export_form_includes_separate_location_fields(self):
        form = AttendanceExportForm()
        choices = {value for value, _ in form.fields["selected_fields"].choices}

        self.assertIn("check_in_location", choices)
        self.assertIn("check_in_maps", choices)
        self.assertIn("check_out_location", choices)
        self.assertIn("check_out_maps", choices)
        self.assertIn("location", choices)
        self.assertIn("maps", choices)

    def test_activity_export_form_includes_separate_location_fields(self):
        form = AttendanceActivityExportForm()
        choices = {value for value, _ in form.fields["selected_fields"].choices}

        self.assertIn("daily_clock_in_location", choices)
        self.assertIn("daily_clock_out_location", choices)
        self.assertIn("daily_break_in_location", choices)
        self.assertIn("daily_break_out_location", choices)
        self.assertIn("daily_lunch_in_location", choices)
        self.assertIn("daily_lunch_out_location", choices)
        self.assertNotIn("clock_in_gps_address", choices)
        self.assertNotIn("clock_in_maps_url", choices)
        self.assertNotIn("clock_out_gps_address", choices)
        self.assertNotIn("clock_out_maps_url", choices)
        self.assertNotIn("gps_address", choices)
        self.assertNotIn("maps_url", choices)
